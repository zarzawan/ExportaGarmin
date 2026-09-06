"""Presentación Excel estática del modelo semántico."""
import copy, json, math, os, re, secrets, unicodedata, warnings
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional
from semantic_common import _number, _parse_date, _sport_family, _date_range
from xlsx_constants import XLSX_MAX_ACTIVITY_SERIES_SAMPLES, XLSX_PRESENTATION_VERSION, XLSX_TRANSLATIONS, XLSX_QUALITY_NAMES

def _flatten_dict(value: Any, prefix: str = "") -> dict[str, Any]:
    result: dict[str, Any] = {}
    if not isinstance(value, dict):
        return {prefix or "value": value}
    for key, nested in value.items():
        path = f"{prefix}.{key}" if prefix else str(key)
        if isinstance(nested, dict):
            result.update(_flatten_dict(nested, path))
        elif isinstance(nested, list):
            result[path] = json.dumps(nested, ensure_ascii=False, default=str)
        else:
            result[path] = nested
    return result


def _excel_safe(value: Any):
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value if math.isfinite(value) else None
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    value = str(value)
    if value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value[:32767]


def _write_table(
    sheet,
    rows: list[dict],
    headers: Optional[list[str]] = None,
    *,
    header_font=None,
    header_fill=None,
):
    if headers is None:
        headers = []
        for row in rows:
            for key in row:
                if key not in headers:
                    headers.append(key)
    if not headers:
        headers = ["sin_datos"]
        rows = [{"sin_datos": "No hay datos disponibles."}]
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.utils import get_column_letter

    preview_rows: list[list[Any]] = []
    widths = [len(str(header)) for header in headers]
    for row in rows[:200]:
        safe_row = [_excel_safe(row.get(header)) for header in headers]
        preview_rows.append(safe_row)
        for index, value in enumerate(safe_row):
            if value is not None:
                widths[index] = max(widths[index], len(str(value)))

    sheet.freeze_panes = "A2"
    last_column = get_column_letter(len(headers))
    sheet.auto_filter.ref = f"A1:{last_column}{len(rows) + 1}"
    for column_number, width in enumerate(widths, start=1):
        letter = get_column_letter(column_number)
        sheet.column_dimensions[letter].width = min(max(width + 2, 10), 45)

    if sheet.parent.write_only:
        header_cells = []
        for header in headers:
            cell = WriteOnlyCell(sheet, value=header)
            if header_font is not None:
                cell.font = header_font
            if header_fill is not None:
                cell.fill = header_fill
            header_cells.append(cell)
        sheet.append(header_cells)
    else:
        sheet.append(headers)
        for cell in sheet[1]:
            if header_font is not None:
                cell.font = header_font
            if header_fill is not None:
                cell.fill = header_fill

    for index, row in enumerate(rows):
        safe_row = (
            preview_rows[index]
            if index < len(preview_rows)
            else [_excel_safe(row.get(header)) for header in headers]
        )
        sheet.append(safe_row)
    return headers


def _xlsx_series_column_base(descriptor: dict, index: int) -> str:
    """Crea un nombre de columna legible y estable para una métrica temporal."""
    raw_name = (
        descriptor.get("field")
        or descriptor.get("source_field")
        or f"valor_{index + 1}"
    )
    plain = unicodedata.normalize("NFKD", str(raw_name))
    plain = "".join(
        character for character in plain
        if not unicodedata.combining(character)
    )
    plain = re.sub(r"[^A-Za-z0-9]+", "_", plain).strip("_").lower()
    if not plain:
        plain = f"valor_{index + 1}"
    if plain[0].isdigit():
        plain = f"metrica_{plain}"
    if plain in {"activity_ref", "sample_index"}:
        plain = f"metrica_{plain}"
    return plain


def _prepare_xlsx_activity_series(activities: list[dict]):
    """Alinea descriptores y muestras sin eliminar los valores nulos posicionales."""
    descriptor_rows: list[dict] = []
    prepared_activities: list[dict] = []
    global_columns: list[str] = []
    signature_columns: dict[tuple, str] = {}
    used_columns = {"activity_ref", "sample_index"}

    for activity in activities:
        series = activity.get("activity_series") or {}
        descriptors = series.get("metric_descriptors") or []
        samples = series.get("samples") or []
        if not isinstance(descriptors, list) or not isinstance(samples, list):
            continue

        activity_ref = activity.get("activity_ref")
        local_columns: list[str] = []
        signature_occurrences: dict[tuple, int] = {}
        for index, raw_descriptor in enumerate(descriptors):
            descriptor = raw_descriptor if isinstance(raw_descriptor, dict) else {}
            canonical = (
                str(descriptor.get("field") or ""),
                str(descriptor.get("source_field") or ""),
                str(descriptor.get("source_unit") or ""),
                json.dumps(
                    descriptor.get("source_factor"),
                    ensure_ascii=False,
                    sort_keys=True,
                    default=str,
                ),
            )
            occurrence = signature_occurrences.get(canonical, 0) + 1
            signature_occurrences[canonical] = occurrence
            signature = (*canonical, occurrence)

            column_name = signature_columns.get(signature)
            if column_name is None:
                base = _xlsx_series_column_base(descriptor, index)
                column_name = base
                suffix = 2
                while column_name in used_columns:
                    column_name = f"{base}_{suffix}"
                    suffix += 1
                signature_columns[signature] = column_name
                used_columns.add(column_name)
                global_columns.append(column_name)

            local_columns.append(column_name)
            descriptor_rows.append({
                "activity_ref": activity_ref,
                "descriptor_index": index,
                "column_name": column_name,
                "field": descriptor.get("field"),
                "source_field": descriptor.get("source_field"),
                "source_unit": descriptor.get("source_unit"),
                "source_factor": descriptor.get("source_factor"),
            })

        if local_columns and samples:
            valid_sample_count = sum(
                1 for sample in samples if isinstance(sample, list)
            )
            prepared_activities.append({
                "activity_ref": activity_ref,
                "columns": local_columns,
                "samples": samples,
                "sample_count": valid_sample_count,
            })

    return descriptor_rows, prepared_activities, global_columns


def _section(model: dict, name: str, default):
    value = model.get(name, default)
    if isinstance(value, dict) and len(value) == 1:
        only_key = next(iter(value))
        if only_key == name:
            return value[only_key]
    return value


def render_xlsx(model: dict, path: Path) -> None:
    from training_analysis import (SCHEMA_VERSION, _DAILY_METRICS, _coverage,
        _sum_week_field, build_period_summary, build_weekly_timeline)
    """Crea el informe Excel deportivo en español sin alterar el modelo TXT."""
    try:
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.chart import BarChart, LineChart, PieChart, Reference
        from openpyxl.chart.axis import DateAxis
        from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.worksheet.table import Table, TableStyleInfo
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError(
            "Falta openpyxl. Ejecuta de nuevo Instalar.bat para instalar "
            "las dependencias de Excel."
        ) from exc

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    metadata = _section(model, "export_metadata", {}) or {}
    race_context = _section(model, "race_context", {}) or {}
    summary = _section(model, "period_summary", {}) or {}
    race_analysis = _section(model, "race_analysis", {}) or {}
    weeks = _section(model, "weekly_timeline", []) or []
    days = _section(model, "daily_health", []) or []
    activities = _section(model, "activities", []) or []
    gear = _section(model, "gear", []) or []
    journal = _section(model, "journal", []) or []
    blood_pressure = _section(model, "blood_pressure", []) or []
    composition = _section(model, "body_composition", []) or []
    quality = _section(model, "data_quality", {}) or {}
    profile = _section(model, "profile", {}) or {}

    # Garmin utiliza valores negativos como centinelas en algunas respuestas.
    # La copia visible evita tratarlos como mediciones sin modificar el modelo
    # semántico que alimenta el TXT ni las hojas técnicas de trazabilidad.
    visible_days = copy.deepcopy(days)
    non_negative_daily_fields = {
        "steps", "distance_m", "active_calories_kcal", "total_calories_kcal",
        "moderate_intensity_minutes", "vigorous_intensity_minutes",
        "resting_heart_rate_bpm", "average_stress", "maximum_stress",
        "body_battery_high", "body_battery_low", "body_battery_charged",
        "body_battery_drained", "average_spo2_pct", "lowest_spo2_pct",
        "average_waking_respiration_brpm", "average_sleep_respiration_brpm",
    }
    non_negative_sleep_fields = {
        "total_sleep_s", "awake_s", "light_sleep_s", "deep_sleep_s",
        "rem_sleep_s", "nap_time_s", "sleep_score",
        "average_sleep_heart_rate_bpm", "average_sleep_stress",
        "average_sleep_spo2_pct",
    }
    non_negative_hrv_fields = {
        "overnight_average_ms", "highest_five_min_average_ms",
        "weekly_average_ms", "baseline_balanced_low_ms",
        "baseline_balanced_high_ms",
    }

    def clear_negative_fields(item, fields):
        if not isinstance(item, dict):
            return
        for field in fields:
            value = _number(item.get(field))
            if value is not None and value < 0:
                item[field] = None

    for visible_day in visible_days:
        clear_negative_fields(visible_day, non_negative_daily_fields)
        clear_negative_fields(visible_day.get("sleep"), non_negative_sleep_fields)
        clear_negative_fields(visible_day.get("hrv"), non_negative_hrv_fields)

    visible_quality = copy.deepcopy(quality)
    descriptor_rows, prepared_series, series_columns = (
        _prepare_xlsx_activity_series(activities)
    )
    available_series_samples = sum(
        item.get("sample_count", 0) for item in prepared_series
    )
    omit_series = (
        available_series_samples > XLSX_MAX_ACTIVITY_SERIES_SAMPLES
    )

    workbook = Workbook(write_only=True)
    workbook.calculation.fullCalcOnLoad = False
    workbook.calculation.forceFullCalc = False
    workbook.calculation.calcMode = "manual"

    colors = {
        "navy": "123A5A",
        "blue": "007CC3",
        "light_blue": "EAF4FA",
        "lighter_blue": "F5FAFD",
        "gray": "E7EAED",
        "dark_gray": "5B6570",
        "green": "D9EAD3",
        "orange": "FCE5CD",
        "red": "F4CCCC",
        "white": "FFFFFF",
    }
    thin_gray = Side(style="thin", color="D4D9DD")
    border = Border(bottom=thin_gray)
    title_font = Font(name="Aptos Display", size=20, bold=True, color=colors["navy"])
    section_font = Font(name="Aptos", size=12, bold=True, color=colors["navy"])
    header_font = Font(name="Aptos", size=10, bold=True, color=colors["white"])
    normal_font = Font(name="Aptos", size=10, color="1F2933")
    secondary_font = Font(name="Aptos", size=9, color=colors["dark_gray"])
    header_fill = PatternFill("solid", fgColor=colors["navy"])
    alternate_fill = PatternFill("solid", fgColor=colors["lighter_blue"])
    light_fill = PatternFill("solid", fgColor=colors["light_blue"])
    warning_fill = PatternFill("solid", fgColor=colors["orange"])
    good_fill = PatternFill("solid", fgColor=colors["green"])

    table_names: set[str] = set()
    mapping_rows: list[dict] = []
    untranslated: set[tuple[str, str]] = set()
    normalized_translations = {
        category: {
            str(source).strip().casefold(): translated
            for source, translated in values.items()
        }
        for category, values in XLSX_TRANSLATIONS.items()
    }
    generic_translations = {
        "strong": "Fuerte",
        "weak": "Flojo",
        "normal": "Normal",
        "trail_running": "Carrera por montaña",
        "trail running": "Carrera por montaña",
        "interval": "Intervalo",
        "other": "Otro",
        "uncategorized": "Sin clasificar",
    }

    def nested(value, path_name, default=None):
        current = value
        for part in path_name.split("."):
            if not isinstance(current, dict):
                return default
            current = current.get(part)
        return default if current is None else current

    def number(value):
        return _number(value)

    def excel_date(value):
        if isinstance(value, datetime):
            return value.replace(tzinfo=None)
        if isinstance(value, date):
            return value
        if not isinstance(value, str) or not value.strip():
            return None
        try:
            return date.fromisoformat(value.strip()[:10])
        except ValueError:
            return None

    def excel_datetime(value):
        if isinstance(value, datetime):
            return value.replace(tzinfo=None)
        if isinstance(value, date):
            return datetime.combine(value, datetime.min.time())
        if not isinstance(value, str) or not value.strip():
            return None
        candidate = value.strip().replace("Z", "+00:00")
        try:
            return datetime.fromisoformat(candidate).replace(tzinfo=None)
        except ValueError:
            parsed = excel_date(candidate)
            return (
                datetime.combine(parsed, datetime.min.time())
                if parsed is not None
                else None
            )

    def excel_time(seconds):
        value = number(seconds)
        return value / 86400.0 if value is not None else None

    def excel_km(metres):
        value = number(metres)
        return value / 1000.0 if value is not None else None

    def excel_pct(value):
        value = number(value)
        return value / 100.0 if value is not None else None

    def yes_no(value):
        if value is None:
            return None
        return "Sí" if bool(value) else "No"

    def humanize(value):
        text = str(value or "").strip()
        if not text:
            return None
        text = re.sub(r"[_\-.]+", " ", text)
        text = re.sub(r"\s+", " ", text).strip()
        return text[:1].upper() + text[1:].lower()

    def translate(category, value):
        if value in (None, ""):
            return None
        text = str(value).strip()
        if text.casefold() in {"none", "null", "n/a", "na"}:
            return None
        translated = normalized_translations.get(category, {}).get(
            text.casefold()
        ) or generic_translations.get(text.casefold())
        if translated is None:
            untranslated.add((category, text))
            return humanize(text)
        return translated

    def translate_known_or_original(category, value):
        """Traduce códigos Garmin sin alterar texto libre de la persona."""
        if value in (None, ""):
            return None
        text = str(value).strip()
        if text.casefold() in {"none", "null", "n/a", "na"}:
            return None
        return normalized_translations.get(category, {}).get(
            text.casefold(), text
        )

    def user_text(value):
        if value in (None, ""):
            return None
        return str(value)

    def narrative(value):
        """Convierte avisos estructurados en prosa sin mostrar JSON."""
        if value in (None, "", [], {}):
            return None
        if isinstance(value, dict):
            if value.get("message"):
                return narrative(value["message"])
            parts = []
            for key, nested_value in value.items():
                rendered = narrative(nested_value)
                if rendered:
                    friendly_key = {
                        "get_sleep_data": "Datos de sueño",
                        "get_hrv_data": "Datos de VFC",
                        "dailySleepDTO": "Datos diarios de sueño",
                        "hrvSummary": "Resumen de VFC",
                        "sleepNeed.actual": "Necesidad de sueño",
                        "recoveryTime": "Tiempo de recuperación",
                        "hrvWeeklyAverage": "Media semanal de VFC",
                        "weather.temp": "Temperatura meteorológica",
                        "directWorkoutRpe": "Esfuerzo percibido registrado",
                        "perceived_exertion_1_10": "Esfuerzo percibido",
                        "summaryDTO": "Resumen de actividad",
                        "lapDTOs": "Vueltas de la actividad",
                        "speed_raw": "Velocidad registrada",
                        "gear_type": "Tipo de equipamiento",
                        "lap_type": "Tipo de vuelta",
                        "session_type": "Tipo de sesión",
                    }.get(str(key), humanize(key))
                    parts.append(f"{friendly_key}: {rendered}")
            return "; ".join(parts) or None
        if isinstance(value, list):
            return "; ".join(
                rendered
                for item in value
                if (rendered := narrative(item))
            ) or None
        if isinstance(value, bool):
            return yes_no(value)
        text_value = str(value)
        if text_value == "Consulta Data Quality.coverage antes de interpretar promedios.":
            return "Revisa la cobertura de datos antes de interpretar promedios."

        # Los mensajes originales permanecen íntegros en TÉCNICO - MODELO. En
        # la hoja visible se agrupan por su significado para que una persona no
        # tenga que interpretar nombres de endpoints, DTO o rutas JSON.
        technical_text = text_value.casefold()
        explanations = []

        def explain_when(tokens, explanation):
            if any(token.casefold() in technical_text for token in tokens):
                if explanation not in explanations:
                    explanations.append(explanation)

        explain_when(
            ("epochs", "sleep raw", "sleep timestamp"),
            "Se convirtieron las marcas de tiempo del sueño a fecha y hora local.",
        )
        explain_when(
            ("lactate-threshold", "lactate_threshold", "lactate threshold"),
            "La velocidad del umbral de lactato se convirtió a metros por segundo.",
        )
        explain_when(
            ("splits.vueltas", "typed_splits.splits", "lapdtos"),
            "Se evitó repetir información de vueltas procedente de distintas fuentes de Garmin.",
        )
        explain_when(
            ("summarydto", "detail", "summary/detail"),
            "Se unificaron los resúmenes generales de cada actividad.",
        )
        explain_when(
            ("split_summaries", "splitsummary", "split summaries"),
            "Se eliminaron resúmenes de intervalos duplicados.",
        )
        explain_when(
            ("totalascent", "totaldescent"),
            "El ascenso y descenso acumulados se interpretaron en metros.",
        )
        explain_when(
            ("speed_raw", "raw speed"),
            "Las velocidades deportivas se conservaron con unidades claras.",
        )
        explain_when(
            ("endpoint",),
            "Algunos datos de Garmin no estuvieron disponibles durante la descarga.",
        )
        if explanations:
            friendly = " ".join(explanations)
            date_range = re.search(
                r"(\d{4}-\d{2}-\d{2})(?:\s*(?:a|/|\.\.|—|-)\s*"
                r"(\d{4}-\d{2}-\d{2}))?",
                text_value,
            )
            if date_range:
                first = excel_date(date_range.group(1))
                last = excel_date(date_range.group(2)) if date_range.group(2) else None
                if first:
                    period = first.strftime("%d/%m/%Y")
                    if last:
                        period += f"–{last.strftime('%d/%m/%Y')}"
                    friendly += f" Periodo: {period}."
            return friendly

        replacements = {
            "get_sleep_data": "datos de sueño",
            "get_hrv_data": "datos de VFC",
            "dailySleepDTO": "datos diarios de sueño",
            "hrvSummary": "resumen de VFC",
            "sleepNeed.actual": "necesidad de sueño",
            "recoveryTime": "tiempo de recuperación",
            "hrvWeeklyAverage": "media semanal de VFC",
            "weather.temp": "temperatura meteorológica",
            "directWorkoutRpe": "esfuerzo percibido registrado",
            "perceived_exertion_1_10": "esfuerzo percibido",
            "summaryDTO": "resumen de actividad",
            "lapDTOs": "vueltas de la actividad",
            "speed_raw": "velocidad registrada",
            "gear_type": "tipo de equipamiento",
            "lap_type": "tipo de vuelta",
            "session_type": "tipo de sesión",
            "raw": "dato original",
            "epochs": "marcas de tiempo",
            "endpoint": "origen de datos",
            "lactate-threshold": "umbral de lactato",
            "splits.vueltas": "vueltas",
            "typed_splits.splits": "vueltas",
            "summary": "resumen",
            "detail": "detalle deportivo",
            "split_summaries": "resúmenes de intervalos",
            "totalAscent": "ascenso acumulado",
            "totalDescent": "descenso acumulado",
        }
        for technical, friendly in replacements.items():
            text_value = re.sub(
                rf"\b{re.escape(technical)}\b",
                friendly,
                text_value,
                flags=re.IGNORECASE,
            )

        def spanish_date(match):
            parsed = excel_date(match.group(0))
            return parsed.strftime("%d/%m/%Y") if parsed else match.group(0)

        text_value = re.sub(r"\b\d{4}-\d{2}-\d{2}\b", spanish_date, text_value)
        text_value = re.sub(
            r"(\d{2}/\d{2}/\d{4})\s*(?:a|/|\.\.|—|-)\s*(\d{2}/\d{2}/\d{4})",
            r"\1–\2",
            text_value,
        )
        return text_value

    def column(
        key,
        header,
        *,
        fmt="General",
        width=15,
        hidden=False,
        description=None,
        source=None,
        conversion="Sin conversión",
        source_unit="—",
        shown_unit="—",
        group=None,
    ):
        return {
            "key": key,
            "header": header,
            "format": fmt,
            "width": width,
            "hidden": hidden,
            "description": description or f"Valor deportivo de {header.lower()}.",
            "source": source or key,
            "conversion": conversion,
            "source_unit": source_unit,
            "shown_unit": shown_unit,
            "group": group,
        }

    def unique_table_name(base):
        plain = unicodedata.normalize("NFKD", base)
        plain = "".join(
            character for character in plain
            if not unicodedata.combining(character)
        )
        plain = re.sub(r"[^A-Za-z0-9]", "", plain)
        candidate = f"Tabla{plain}"[:250] or "TablaDatos"
        suffix = 2
        while candidate in table_names:
            candidate = f"Tabla{plain}{suffix}"[:250]
            suffix += 1
        table_names.add(candidate)
        return candidate

    def register_mapping(sheet_name, columns):
        for item in columns:
            mapping_rows.append({
                "sheet": sheet_name,
                "visible_name": item["header"],
                "internal_field": item["source"],
                "conversion": item["conversion"],
                "source_unit": item["source_unit"],
                "shown_unit": item["shown_unit"],
                "description": item["description"],
            })

    def append_table(sheet, rows, columns, table_base, start_row=1):
        headers = [item["header"] for item in columns]
        # Una altura fija evita que Excel expanda la cabecera por columnas
        # avanzadas ocultas. El ajuste permite dos líneas en los títulos largos.
        sheet.row_dimensions[start_row].height = 33
        for item in columns:
            letter = get_column_letter(columns.index(item) + 1)
            sheet.column_dimensions[letter].width = item["width"]
            sheet.column_dimensions[letter].hidden = item["hidden"]
            sheet.column_dimensions[letter].outlineLevel = item.get(
                "outline_level", 0
            )

        header_cells = []
        for item in columns:
            cell = WriteOnlyCell(sheet, value=item["header"])
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            header_cells.append(cell)
        sheet.append(header_cells)

        for row_index, row in enumerate(rows, 1):
            cells = []
            for item in columns:
                cell = WriteOnlyCell(sheet, value=_excel_safe(row.get(item["key"])))
                cell.font = normal_font
                cell.number_format = (row.get("__formats__") or {}).get(
                    item["key"], item["format"]
                )
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=item["width"] >= 28,
                )
                cell.border = border
                if row_index % 2 == 0:
                    cell.fill = alternate_fill
                cells.append(cell)
            sheet.append(cells)

        last_column = get_column_letter(len(columns))
        last_row = start_row + len(rows)
        table = Table(
            displayName=unique_table_name(table_base),
            ref=f"A{start_row}:{last_column}{last_row}",
        )
        table._initialise_columns()
        for table_column, header in zip(table.tableColumns, headers):
            table_column.name = header
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        with warnings.catch_warnings():
            warnings.filterwarnings(
                "ignore",
                message="In write-only mode you must add table columns manually",
                category=UserWarning,
            )
            sheet.add_table(table)
        return last_row

    def add_table_sheet(
        name,
        rows,
        columns,
        *,
        freeze="A2",
        hidden=False,
        primary_keys=None,
    ):
        if not rows:
            return None
        sheet = workbook.create_sheet(name)
        sheet.freeze_panes = freeze
        sheet.sheet_view.showGridLines = True
        rendered_columns = columns
        if primary_keys and not hidden:
            primary_keys = set(primary_keys)
            sheet.sheet_properties.outlinePr.summaryRight = True
            rendered_columns = []
            for item in columns:
                rendered = dict(item)
                if item["key"] not in primary_keys and not item["hidden"]:
                    rendered["hidden"] = True
                    rendered["outline_level"] = 1
                rendered_columns.append(rendered)
        append_table(sheet, rows, rendered_columns, name, 1)
        if hidden:
            sheet.sheet_state = "hidden"
        else:
            register_mapping(name, rendered_columns)
        return sheet

    def add_mapping_entry(sheet_name, visible, source, conversion, unit, description):
        mapping_rows.append({
            "sheet": sheet_name,
            "visible_name": visible,
            "internal_field": source,
            "conversion": conversion,
            "source_unit": unit,
            "shown_unit": unit,
            "description": description,
        })

    activity_by_ref = {
        str(item.get("activity_ref")): item
        for item in activities
        if item.get("activity_ref")
    }
    gear_by_ref = {
        str(item.get("gear_ref")): item
        for item in gear
        if item.get("gear_ref")
    }

    def activity_label(reference):
        activity = activity_by_ref.get(str(reference), {})
        return user_text(activity.get("name")) or "Actividad sin nombre"

    def gear_label(item):
        name = user_text(item.get("gear_name"))
        manufacturer = user_text(item.get("manufacturer"))
        model_name = user_text(item.get("model"))
        parts = []
        if name:
            parts.append(name)
        model_label = " ".join(part for part in (manufacturer, model_name) if part)
        if model_label and model_label.casefold() not in " ".join(parts).casefold():
            parts.append(model_label)
        return " — ".join(parts) or "Equipamiento sin nombre"

    def activity_gear_label(activity):
        labels = [
            gear_label(gear_by_ref[str(reference)])
            for reference in activity.get("gear_refs", []) or []
            if str(reference) in gear_by_ref
        ]
        return " · ".join(labels) or None

    def is_microactivity(activity):
        """Detecta registros deportivos demasiado breves para indicadores."""
        if _sport_family(activity.get("sport")) not in {"running", "cycling"}:
            return False
        duration = number(activity.get("duration_s"))
        distance = number(activity.get("distance_m"))
        return (
            duration is not None
            and distance is not None
            and duration < 60
            and distance < 100
        )

    def visible_session_type(activity):
        classification = activity.get("classification") or {}
        raw_type = str(classification.get("type") or "unknown").strip()
        normalized = raw_type.casefold()
        evidence = {
            str(item).casefold()
            for item in classification.get("evidence", []) or []
        }

        # El tipo previsto por la persona es la fuente de mayor prioridad.
        if (
            classification.get("source") == "user_provided"
            or "manual_intended_session_type" in evidence
        ):
            return translate("session_type", raw_type) or "Sin clasificar"

        def folded(value):
            plain = unicodedata.normalize("NFKD", str(value or ""))
            plain = "".join(
                character for character in plain
                if not unicodedata.combining(character)
            )
            return re.sub(r"\s+", " ", plain.casefold()).strip()

        name = folded(activity.get("name"))
        event_type = folded(activity.get("garmin_event_type"))
        if "carrera facil" in name:
            return "Rodaje fácil"
        if "carrera larga" in name or "tirada larga" in name:
            return "Tirada larga"

        competition_name = any(token in name for token in (
            "maraton", "marathon", "media maraton", "half marathon",
            "1/2 maraton", "competicion",
        ))
        if (
            normalized == "race"
            or "race" in event_type
            or "competition" in event_type
            or competition_name
        ):
            return "Competición"
        if "tempo" in name:
            return "Tempo"
        if any(token in name for token in (
            "intervalos", "series", "repeticiones",
        )):
            return "Intervalos"

        # Solo se aceptan metadatos nominales explícitos del entrenamiento. La
        # mera existencia de interval_summaries, vueltas automáticas,
        # INTERVAL_ACTIVE o RWD_RUN no constituye evidencia.
        structured_text = " ".join(
            folded(activity.get(field))
            for field in (
                "workout_name", "structured_workout_name", "workout_title",
                "training_plan_name", "workout_type",
            )
            if activity.get(field)
        )
        if any(token in structured_text for token in (
            "interval", "series", "repeticion", "repeat",
        )):
            return "Intervalos"
        if "tempo" in structured_text:
            return "Tempo"

        step_types = [
            folded(lap.get("step_type"))
            for lap in activity.get("laps", []) or []
            if isinstance(lap, dict)
        ]
        work_steps = sum(
            any(token in step for token in ("work", "repeat", "repeticion"))
            for step in step_types
        )
        recovery_steps = sum(
            any(token in step for token in ("recovery", "rest", "recuperacion"))
            for step in step_types
        )
        if work_steps >= 2 and recovery_steps >= 1:
            return "Intervalos"

        # Las reglas conservadoras ya existentes para rodaje fácil, tirada
        # larga, fuerza y entrenamiento cruzado siguen siendo válidas. La
        # clasificación interval derivada de structured_lap_steps se descarta.
        if normalized == "interval":
            return "Sin clasificar"
        if normalized in {"", "none", "unknown", "uncategorized"}:
            return "Sin clasificar"
        return translate("session_type", raw_type)

    def interval_role(interval):
        raw_type = str(interval.get("interval_type") or "").strip().casefold()
        if raw_type in {"interval_warmup", "warmup"}:
            return "warmup"
        if raw_type in {"interval_cooldown", "cooldown"}:
            return "cooldown"
        if raw_type in {"interval_recovery", "interval_rest", "recovery"}:
            return "recovery"
        if raw_type in {"interval_active", "interval", "work", "repeat"}:
            return "work"
        if raw_type == "rest":
            return "pause"
        if raw_type in {"rwd_run", "run", "active"}:
            return "total_candidate"
        return "automatic_or_generic"

    def interval_values_are_close(first, second):
        """Compara resúmenes tolerando pausas breves y redondeos de Garmin."""
        first_count = number(first.get("interval_count"))
        second_count = number(second.get("interval_count"))
        if (
            first_count is not None
            and second_count is not None
            and first_count != second_count
        ):
            return False
        for field, absolute_tolerance in (
            ("distance_m", 100.0),
            ("duration_s", 60.0),
        ):
            left = number(first.get(field))
            right = number(second.get(field))
            if left is None or right is None:
                continue
            tolerance = max(
                absolute_tolerance,
                max(abs(left), abs(right)) * 0.02,
            )
            if abs(left - right) > tolerance:
                return False
        return True

    def visible_interval_items(activity):
        """Devuelve solo bloques útiles para un corredor o entrenador."""
        intervals = [
            item for item in activity.get("interval_summaries", []) or []
            if isinstance(item, dict)
        ]
        if not intervals:
            return []

        session_type = visible_session_type(activity)
        roles = [interval_role(item) for item in intervals]
        structured = (
            session_type in {"Intervalos", "Tempo"}
            or "warmup" in roles
            or "cooldown" in roles
            or ("work" in roles and "recovery" in roles)
        )
        if not structured:
            return []

        labels = {
            "warmup": ("Calentamiento", "Bloque planificado"),
            "work": ("Trabajo activo", "Resumen de bloques"),
            "recovery": ("Recuperación", "Bloque planificado"),
            "pause": ("Pausa", "Bloque planificado"),
            "cooldown": ("Enfriamiento", "Bloque planificado"),
            "total": ("Total de actividad", "Total de actividad"),
        }
        selected = []
        for interval in intervals:
            role = interval_role(interval)
            if role not in {"warmup", "work", "recovery", "pause", "cooldown"}:
                continue
            if role == "pause":
                duration = number(interval.get("duration_s")) or 0.0
                count = number(interval.get("interval_count")) or 0.0
                if duration < 30 and count < 1:
                    continue
            if any(
                existing[0] == role
                and interval_values_are_close(existing[1], interval)
                for existing in selected
            ):
                continue
            selected.append((role, interval))

        # Los totales ya visibles en ACTIVIDADES se omiten. Solo se conserva
        # uno cuando a la actividad le falta su resumen general y Garmin sí lo
        # proporciona, con una prioridad explícita y siempre al final.
        activity_has_total = (
            number(activity.get("distance_m")) is not None
            or number(activity.get("duration_s")) is not None
        )
        if not activity_has_total:
            total_priority = {"run": 0, "rwd_run": 1, "active": 2}
            total_candidates = sorted(
                (
                    item for item in intervals
                    if interval_role(item) == "total_candidate"
                ),
                key=lambda item: total_priority.get(
                    str(item.get("interval_type") or "").strip().casefold(),
                    99,
                ),
            )
            if total_candidates:
                selected.append(("total", total_candidates[0]))

        return [
            (interval, *labels[role])
            for role, interval in selected
        ]

    def interval_pace(value, interval):
        distance = number(interval.get("distance_m"))
        moving = number(interval.get("moving_duration_s"))
        pace = number(value)
        if (
            pace is None
            or distance is None
            or distance < 100
            or (moving is not None and moving <= 0)
            or pace > 3600
        ):
            return None
        return excel_time(pace)

    def lap_pace(value, lap):
        distance = number(lap.get("distance_m"))
        moving = number(lap.get("moving_duration_s"))
        pace = number(value)
        if (
            pace is None
            or distance is None
            or distance < 100
            or (moving is not None and moving <= 0)
            or pace > 3600
        ):
            return None
        return excel_time(pace)

    microactivities = [item for item in activities if is_microactivity(item)]
    coach_activities = [item for item in activities if not is_microactivity(item)]
    valid_evaluated = sum(
        1 for item in coach_activities if item.get("self_evaluation")
    )
    visible_quality.setdefault("coverage", {})["self_evaluation"] = {
        "available_activities": valid_evaluated,
        "expected_activities": len(coach_activities),
        "missing_activities": len(coach_activities) - valid_evaluated,
        "coverage_pct": (
            round(valid_evaluated * 100.0 / len(coach_activities), 1)
            if coach_activities
            else None
        ),
    }
    period_start = _parse_date(
        nested(summary, "period.start_date") or metadata.get("start_date")
    )
    period_end = _parse_date(
        nested(summary, "period.end_date") or metadata.get("end_date")
    )
    if period_start and period_end:
        visible_quality.setdefault("coverage", {})["average_stress"] = _coverage(
            visible_days,
            _date_range(period_start, period_end),
            _DAILY_METRICS["average_stress"],
        )
        coach_weeks = build_weekly_timeline(
            coach_activities,
            visible_days,
            period_start,
            period_end,
            reference_date=period_end,
        )
        coach_summary = build_period_summary(
            coach_activities,
            visible_days,
            period_start,
            period_end,
            coach_weeks,
        )
    else:
        coach_weeks = weeks
        coach_summary = summary

    def compare_complete_week_blocks(rows):
        complete = [row for row in rows if row.get("status") == "complete"]
        block_size = min(4, len(complete) // 2)
        if block_size < 1:
            return {
                "status": "insufficient_data",
                "complete_weeks_available": len(complete),
                "minimum_required": 2,
            }
        previous = complete[-2 * block_size:-block_size]
        recent = complete[-block_size:]
        metrics = {}
        for field in (
            "running_distance_m",
            "running_duration_s",
            "running_sessions",
            "longest_run_distance_m",
            "garmin_training_load_total",
            "session_rpe_load_total",
        ):
            previous_value = _sum_week_field(previous, field)
            recent_value = _sum_week_field(recent, field)
            delta = recent_value - previous_value
            metrics[field] = {
                "previous_block_total": round(previous_value, 1),
                "recent_block_total": round(recent_value, 1),
                "percentage_change": (
                    round(delta * 100.0 / previous_value, 1)
                    if previous_value else None
                ),
            }
        return {
            "status": "available",
            "block_size": block_size,
            "complete_weeks_available": len(complete),
            "metrics": metrics,
        }

    def coverage_mean(week, metric):
        value = week.get(metric)
        return value.get("mean") if isinstance(value, dict) else None

    def coverage_pct(week, metric):
        value = week.get(metric)
        return excel_pct(value.get("coverage_pct")) if isinstance(value, dict) else None

    def zone_percentage(week, zone_number):
        distribution = nested(week, "heart_rate_distribution.zones", []) or []
        for zone in distribution:
            if number(zone.get("zone")) == zone_number:
                return excel_pct(zone.get("percentage_of_classified_and_below"))
        return None

    week_rows = []
    for week in coach_weeks:
        week_rows.append({
            "week": week.get("iso_week"),
            "from": excel_date(week.get("scope_start_date") or week.get("week_start_date")),
            "to": excel_date(week.get("scope_end_date") or week.get("week_end_date")),
            "status": translate(
                "week_status",
                "empty" if not number(week.get("training_sessions_total")) else week.get("status"),
            ),
            "sessions": week.get("training_sessions_total"),
            "training_days": week.get("days_with_any_training"),
            "rest_days": week.get("days_without_recorded_training"),
            "run_sessions": week.get("running_sessions"),
            "run_km": excel_km(week.get("running_distance_m")),
            "run_time": excel_time(week.get("running_duration_s")),
            "longest_km": excel_km(week.get("longest_run_distance_m")),
            "bike_sessions": week.get("cycling_sessions"),
            "bike_km": excel_km(week.get("cycling_distance_m")),
            "bike_time": excel_time(week.get("cycling_duration_s")),
            "strength_sessions": week.get("strength_sessions"),
            "strength_time": excel_time(week.get("strength_duration_s")),
            "garmin_load": week.get("garmin_training_load_total"),
            "rpe_load": week.get("session_rpe_load_total"),
            "average_rpe": week.get("average_rpe_1_10"),
            "sleep": excel_time(coverage_mean(week, "sleep_duration_s")),
            "sleep_score": coverage_mean(week, "sleep_score"),
            "hrv": coverage_mean(week, "hrv_overnight_ms"),
            "resting_hr": coverage_mean(week, "resting_heart_rate_bpm"),
            "stress": coverage_mean(week, "average_stress"),
            "body_battery": coverage_mean(week, "body_battery_high"),
            "sleep_coverage": coverage_pct(week, "sleep_duration_s"),
            "evaluation_coverage": excel_pct(nested(week, "self_evaluation_coverage.coverage_pct")),
            **{f"zone_{index}": zone_percentage(week, index) for index in range(1, 6)},
        })

    activity_rows = []
    interval_rows = []
    lap_rows = []
    zone_rows = []
    for activity in activities:
        reference = activity.get("activity_ref")
        activity_date = excel_date(activity.get("date"))
        activity_name = user_text(activity.get("name")) or "Actividad sin nombre"
        activity_rows.append({
            "date": activity_date,
            "time_bucket": translate("time_bucket", activity.get("start_time_bucket")),
            "name": activity_name,
            "sport": translate("sport", activity.get("sport")),
            "session_type": visible_session_type(activity),
            "record_status": (
                "Registro muy breve" if is_microactivity(activity) else "Válido para análisis"
            ),
            "distance_km": excel_km(activity.get("distance_m")),
            "duration": excel_time(activity.get("duration_s")),
            "moving_time": excel_time(activity.get("moving_duration_s")),
            "average_pace": excel_time(activity.get("average_pace_s_per_km")),
            "best_pace": excel_time(activity.get("best_pace_s_per_km")),
            "average_hr": activity.get("average_heart_rate_bpm"),
            "maximum_hr": activity.get("maximum_heart_rate_bpm"),
            "average_power": activity.get("average_power_w"),
            "normalized_power": activity.get("normalized_power_w"),
            "average_cadence": (
                activity.get("average_cadence_spm")
                if activity.get("average_cadence_spm") is not None
                else activity.get("average_cycling_cadence_rpm")
            ),
            "elevation_gain": activity.get("elevation_gain_m"),
            "aerobic_effect": activity.get("aerobic_training_effect"),
            "anaerobic_effect": activity.get("anaerobic_training_effect"),
            "benefit": translate("training_benefit", activity.get("training_effect_label")),
            "garmin_load": activity.get("training_load"),
            "calories": activity.get("calories_kcal"),
            "sweat": activity.get("estimated_sweat_loss_ml"),
            "rpe": nested(activity, "self_evaluation.perceived_exertion_1_10"),
            "feeling": translate("feeling", nested(activity, "self_evaluation.feeling")),
            "temperature": activity.get("average_temperature_c"),
            "gear": activity_gear_label(activity),
            "activity_ref": reference,
        })

        for interval, interval_name, level in visible_interval_items(activity):
            interval_rows.append({
                "date": activity_date,
                "activity": activity_name,
                "type": interval_name,
                "level": level,
                "count": interval.get("interval_count"),
                "distance_km": excel_km(interval.get("distance_m")),
                "duration": excel_time(interval.get("duration_s")),
                "moving_time": excel_time(interval.get("moving_duration_s")),
                "average_pace": interval_pace(interval.get("average_pace_s_per_km"), interval),
                "moving_pace": interval_pace(interval.get("moving_pace_s_per_km"), interval),
                "best_pace": interval_pace(interval.get("best_pace_s_per_km"), interval),
                "average_hr": interval.get("average_heart_rate_bpm"),
                "maximum_hr": interval.get("maximum_heart_rate_bpm"),
                "average_power": interval.get("average_power_w"),
                "maximum_power": interval.get("maximum_power_w"),
                "cadence": interval.get("average_cadence_spm"),
                "calories": interval.get("calories_kcal"),
                "elevation_gain": interval.get("elevation_gain_m"),
                "elevation_loss": interval.get("elevation_loss_m"),
                "gap": interval_pace(interval.get("grade_adjusted_pace_s_per_km"), interval),
                "activity_ref": reference,
            })

        for lap in activity.get("laps", []) or []:
            lap_rows.append({
                "date": activity_date,
                "activity": activity_name,
                "lap_index": lap.get("lap_index"),
                "lap_type": translate("lap_type", lap.get("step_type") or lap.get("lap_type")),
                "distance_m": lap.get("distance_m"),
                "duration": excel_time(lap.get("duration_s")),
                "moving_time": excel_time(lap.get("moving_duration_s")),
                "average_pace": lap_pace(lap.get("average_pace_s_per_km"), lap),
                "best_pace": lap_pace(lap.get("best_pace_s_per_km"), lap),
                "average_hr": lap.get("average_heart_rate_bpm"),
                "maximum_hr": lap.get("maximum_heart_rate_bpm"),
                "average_power": lap.get("average_power_w"),
                "maximum_power": lap.get("maximum_power_w"),
                "normalized_power": lap.get("normalized_power_w"),
                "cadence": lap.get("average_cadence_spm"),
                "stride": lap.get("average_stride_length_cm"),
                "ground_contact": lap.get("average_ground_contact_time_ms"),
                "vertical_oscillation": lap.get("average_vertical_oscillation_cm"),
                "calories": lap.get("calories_kcal"),
                "elevation_gain": lap.get("elevation_gain_m"),
                "elevation_loss": lap.get("elevation_loss_m"),
                "partial": yes_no(lap.get("partial_lap")),
                "activity_ref": reference,
            })

        for zone_type, zones in (
            ("heart_rate", activity.get("hr_zones", [])),
            ("power", activity.get("power_zones", [])),
        ):
            valid_zones = [item for item in zones or [] if isinstance(item, dict)]
            total_duration = sum(number(item.get("duration_s")) or 0 for item in valid_zones)
            for zone_index, zone in enumerate(valid_zones):
                zone_number = number(zone.get("zone"))
                duration = number(zone.get("duration_s"))
                lower_key = "low_boundary_bpm" if zone_type == "heart_rate" else "low_boundary_w"
                lower = number(zone.get(lower_key))
                next_lower = (
                    number(valid_zones[zone_index + 1].get(lower_key))
                    if zone_index + 1 < len(valid_zones)
                    else None
                )
                zone_rows.append({
                    "date": activity_date,
                    "activity": activity_name,
                    "zone_type": translate("zone_type", zone_type),
                    "zone": (
                        translate("zone_type", "below_zone_1")
                        if zone_number == 0
                        else int(zone_number) if zone_number is not None else None
                    ),
                    "duration": excel_time(duration),
                    "percentage": (
                        duration / total_duration
                        if duration is not None and total_duration > 0
                        else None
                    ),
                    "lower": lower,
                    "upper": next_lower,
                    "unit": "lpm" if zone_type == "heart_rate" else "W",
                    "activity_ref": reference,
                })

    daily_rows = []
    habit_rows = []
    for item in visible_days:
        sleep = item.get("sleep") or {}
        hrv = item.get("hrv") or {}
        daily_rows.append({
            "date": excel_date(item.get("date")),
            "steps": item.get("steps"),
            "distance_km": excel_km(item.get("distance_m")),
            "active_calories": item.get("active_calories_kcal"),
            "total_calories": item.get("total_calories_kcal"),
            "moderate_minutes": item.get("moderate_intensity_minutes"),
            "vigorous_minutes": item.get("vigorous_intensity_minutes"),
            "resting_hr": item.get("resting_heart_rate_bpm"),
            "average_stress": item.get("average_stress"),
            "maximum_stress": item.get("maximum_stress"),
            "body_battery_high": item.get("body_battery_high"),
            "body_battery_low": item.get("body_battery_low"),
            "body_battery_charged": item.get("body_battery_charged"),
            "body_battery_drained": item.get("body_battery_drained"),
            "average_spo2": excel_pct(item.get("average_spo2_pct")),
            "lowest_spo2": excel_pct(item.get("lowest_spo2_pct")),
            "waking_respiration": item.get("average_waking_respiration_brpm"),
            "sleep_available": yes_no(sleep.get("valid_sleep")),
            "sleep_start": excel_datetime(sleep.get("sleep_start_local")),
            "sleep_end": excel_datetime(sleep.get("sleep_end_local")),
            "total_sleep": excel_time(sleep.get("total_sleep_s")),
            "awake": excel_time(sleep.get("awake_s")),
            "light_sleep": excel_time(sleep.get("light_sleep_s")),
            "deep_sleep": excel_time(sleep.get("deep_sleep_s")),
            "rem_sleep": excel_time(sleep.get("rem_sleep_s")),
            "naps": excel_time(sleep.get("nap_time_s")),
            "sleep_score": sleep.get("sleep_score"),
            "sleep_rating": translate("sleep", sleep.get("sleep_score_qualifier")),
            "sleep_hr": sleep.get("average_sleep_heart_rate_bpm"),
            "sleep_stress": sleep.get("average_sleep_stress"),
            "sleep_spo2": excel_pct(sleep.get("average_sleep_spo2_pct")),
            "sleep_respiration": item.get("average_sleep_respiration_brpm"),
            "hrv_average": hrv.get("overnight_average_ms"),
            "hrv_best": hrv.get("highest_five_min_average_ms"),
            "hrv_low": hrv.get("baseline_balanced_low_ms"),
            "hrv_high": hrv.get("baseline_balanced_high_ms"),
            "hrv_status": translate("hrv", hrv.get("status")),
            "hrv_week": hrv.get("weekly_average_ms"),
        })
        for habit in item.get("lifestyle_logs", []) or []:
            if not isinstance(habit, dict):
                continue
            habit_rows.append({
                "date": excel_date(habit.get("date") or item.get("date")),
                "habit": translate("habit", habit.get("behaviour")),
                "status": (
                    yes_no(habit.get("status"))
                    if isinstance(habit.get("status"), bool)
                    else humanize(habit.get("status"))
                ),
                "amount": habit.get("amount"),
                "note": user_text(habit.get("note")),
            })

    composition_rows = [{
        "date": excel_date(item.get("date")),
        "weight": item.get("weight_kg"),
        "bmi": item.get("bmi"),
        "fat": excel_pct(item.get("body_fat_pct")),
        "water": excel_pct(item.get("body_water_pct")),
        "muscle": item.get("muscle_mass_kg"),
        "bone": item.get("bone_mass_kg"),
    } for item in composition]
    pressure_rows = [{
        "datetime": excel_datetime(item.get("timestamp")),
        "systolic": item.get("systolic_mmhg"),
        "diastolic": item.get("diastolic_mmhg"),
        "pulse": item.get("pulse_bpm"),
    } for item in blood_pressure]

    gear_usage = {}
    for activity in activities:
        for reference in activity.get("gear_refs", []) or []:
            usage = gear_usage.setdefault(str(reference), {
                "distance_m": 0.0,
                "activities": 0,
                "last_date": None,
            })
            usage["distance_m"] += number(activity.get("distance_m")) or 0.0
            usage["activities"] += 1
            activity_date = excel_date(activity.get("date"))
            if activity_date and (usage["last_date"] is None or activity_date > usage["last_date"]):
                usage["last_date"] = activity_date
    gear_rows = []
    for item in gear:
        reference = str(item.get("gear_ref") or "")
        usage = gear_usage.get(reference, {})
        status = item.get("status")
        if status is None and item.get("retired") is not None:
            status = "retired" if item.get("retired") else "active"
        gear_rows.append({
            "name": user_text(item.get("gear_name")),
            "model": " ".join(
                value for value in (
                    user_text(item.get("manufacturer")),
                    user_text(item.get("model")),
                ) if value
            ) or None,
            "type": translate("gear_type", item.get("type")),
            "status": translate("gear_status", status),
            "total_km": excel_km(item.get("total_distance_m")),
            "period_km": excel_km(usage.get("distance_m")),
            "activities": usage.get("activities"),
            "last_date": usage.get("last_date"),
            "gear_ref": item.get("gear_ref"),
        })

    consolidated_journal = {}
    for item in journal:
        reference = item.get("activity_ref")
        journal_date = excel_date(item.get("date"))
        key = (journal_date, str(reference or ""))
        grouped = consolidated_journal.setdefault(key, {
            "date": journal_date,
            "activity_ref": reference,
            "notes": [],
            "planned_types": [],
        })
        note = user_text(item.get("note"))
        planned_type = translate_known_or_original(
            "session_type", item.get("intended_session_type")
        )
        if note and note not in grouped["notes"]:
            grouped["notes"].append(note)
        if planned_type and planned_type not in grouped["planned_types"]:
            grouped["planned_types"].append(planned_type)

    journal_rows = []
    for grouped in consolidated_journal.values():
        if not grouped["notes"] and not grouped["planned_types"]:
            continue
        reference = grouped["activity_ref"]
        activity_date = excel_date(
            activity_by_ref.get(str(reference), {}).get("date")
        ) if reference else None
        mismatch = (
            grouped["date"] is not None
            and activity_date is not None
            and grouped["date"] != activity_date
        )
        journal_rows.append({
            "journal_date": grouped["date"],
            "activity_date": activity_date,
            "activity": activity_label(reference) if reference else None,
            "note": "\n".join(grouped["notes"]) or None,
            "planned_type": " · ".join(grouped["planned_types"]) or None,
            "warning": (
                "La fecha del diario no coincide con la fecha de la actividad."
                if mismatch else None
            ),
            "activity_ref": reference,
        })

    quality_rows = []
    for metric, values in (visible_quality.get("coverage") or {}).items():
        if not isinstance(values, dict):
            continue
        available = values.get("available_days", values.get("available_activities"))
        expected = values.get("expected_days", values.get("expected_activities"))
        missing = values.get("missing_days", values.get("missing_activities"))
        coverage_value = excel_pct(values.get("coverage_pct"))
        status = None
        if coverage_value is not None:
            status = "Buena" if coverage_value >= 0.8 else "Parcial" if coverage_value >= 0.5 else "Baja"
        mean_value = values.get("mean")
        median_value = values.get("median")
        row_formats = {}
        if metric == "sleep_duration_s":
            mean_value = excel_time(mean_value)
            median_value = excel_time(median_value)
            row_formats = {"mean": "[h]:mm:ss", "median": "[h]:mm:ss"}
        missing_periods = []
        for period_value in values.get("missing_date_ranges") or []:
            if isinstance(period_value, str):
                match = re.fullmatch(
                    r"(\d{4}-\d{2}-\d{2})(?:\s*(?:a|/|\.\.|—|-)\s*(\d{4}-\d{2}-\d{2}))?",
                    period_value.strip(),
                )
                if match:
                    first = excel_date(match.group(1))
                    last = excel_date(match.group(2)) if match.group(2) else None
                    period_value = first.strftime("%d/%m/%Y") if first else period_value
                    if last:
                        period_value += f"–{last.strftime('%d/%m/%Y')}"
            missing_periods.append(str(period_value))
        quality_rows.append({
            "metric": XLSX_QUALITY_NAMES.get(metric, humanize(metric)),
            "mean": mean_value,
            "median": median_value,
            "available": available,
            "expected": expected,
            "missing": missing,
            "coverage": coverage_value,
            "missing_periods": ", ".join(missing_periods) or None,
            "status": status,
            "__formats__": row_formats,
        })

    # Columnas visibles. Las conversiones ya están aplicadas en las filas, por
    # lo que los valores siguen siendo números y fechas reales de Excel.
    week_columns = [
        column("week", "Semana", width=13, source="iso_week", description="Identificador ISO de la semana."),
        column("from", "Desde", fmt="dd/mm/yyyy", width=12, source="scope_start_date", description="Primer día incluido de la semana."),
        column("to", "Hasta", fmt="dd/mm/yyyy", width=12, source="scope_end_date", description="Último día incluido de la semana."),
        column("status", "Estado de la semana", width=19, source="status", description="Indica si la semana está completa, parcial o vacía."),
        column("sessions", "Sesiones totales", width=14, source="training_sessions_total"),
        column("training_days", "Días entrenados", width=14, source="days_with_any_training"),
        column("rest_days", "Días sin entrenamiento", width=19, source="days_without_recorded_training"),
        column("run_sessions", "Sesiones de carrera", width=17, source="running_sessions"),
        column("run_km", "Carrera (km)", fmt="0.00", width=14, source="running_distance_m", conversion="Metros ÷ 1.000", source_unit="m", shown_unit="km"),
        column("run_time", "Tiempo corriendo", fmt="[h]:mm:ss", width=16, source="running_duration_s", conversion="Segundos ÷ 86.400", source_unit="s", shown_unit="duración Excel"),
        column("longest_km", "Tirada más larga (km)", fmt="0.00", width=20, source="longest_run_distance_m", conversion="Metros ÷ 1.000", source_unit="m", shown_unit="km"),
        column("bike_sessions", "Sesiones de ciclismo", width=18, source="cycling_sessions"),
        column("bike_km", "Ciclismo (km)", fmt="0.00", width=14, source="cycling_distance_m", conversion="Metros ÷ 1.000", source_unit="m", shown_unit="km"),
        column("bike_time", "Tiempo de ciclismo", fmt="[h]:mm:ss", width=17, source="cycling_duration_s", conversion="Segundos ÷ 86.400", source_unit="s", shown_unit="duración Excel"),
        column("strength_sessions", "Sesiones de fuerza", width=17, source="strength_sessions"),
        column("strength_time", "Tiempo de fuerza", fmt="[h]:mm:ss", width=16, source="strength_duration_s", conversion="Segundos ÷ 86.400", source_unit="s", shown_unit="duración Excel"),
        column("garmin_load", "Carga Garmin", fmt="0.0", width=14, source="garmin_training_load_total"),
        column("rpe_load", "Carga por esfuerzo", fmt="0.0", width=17, source="session_rpe_load_total"),
        column("average_rpe", "Esfuerzo percibido medio", fmt="0.0", width=22, source="average_rpe_1_10", shown_unit="1–10"),
        column("sleep", "Sueño medio", fmt="[h]:mm:ss", width=14, source="sleep_duration_s.mean", conversion="Segundos ÷ 86.400", source_unit="s", shown_unit="duración Excel"),
        column("sleep_score", "Puntuación de sueño", fmt="0.0", width=19, source="sleep_score.mean"),
        column("hrv", "VFC nocturna (ms)", fmt="0", width=17, source="hrv_overnight_ms.mean", shown_unit="ms"),
        column("resting_hr", "Pulso en reposo (lpm)", fmt="0", width=20, source="resting_heart_rate_bpm.mean", shown_unit="lpm"),
        column("stress", "Estrés medio", fmt="0.0", width=14, source="average_stress.mean"),
        column("body_battery", "Body Battery máximo", fmt="0.0", width=19, source="body_battery_high.mean"),
        column("sleep_coverage", "Cobertura de sueño", fmt="0.0%", width=18, source="sleep_duration_s.coverage_pct", conversion="Porcentaje ÷ 100", source_unit="0–100", shown_unit="%"),
        column("evaluation_coverage", "Cobertura de autoevaluación", fmt="0.0%", width=24, source="self_evaluation_coverage.coverage_pct", conversion="Porcentaje ÷ 100", source_unit="0–100", shown_unit="%"),
        *[
            column(f"zone_{index}", f"Tiempo en zona {index}", fmt="0.0%", width=16, source=f"heart_rate_distribution.zones[{index}].percentage", conversion="Porcentaje ÷ 100", source_unit="0–100", shown_unit="%")
            for index in range(1, 6)
        ],
    ]

    activity_columns = [
        column("date", "Fecha", fmt="dd/mm/yyyy", width=12, source="date"),
        column("time_bucket", "Franja horaria", width=15, source="start_time_bucket"),
        column("name", "Actividad", width=32, source="name", description="Título original de la actividad escrito en Garmin."),
        column("sport", "Deporte", width=20, source="sport"),
        column("session_type", "Tipo de sesión", width=21, source="classification.type"),
        column("record_status", "Estado del registro", width=22, source="xlsx.microactivity_status", description="Indica si el registro se usa en los indicadores del entrenador."),
        column("distance_km", "Distancia (km)", fmt="0.00", width=15, source="distance_m", conversion="Metros ÷ 1.000", source_unit="m", shown_unit="km"),
        column("duration", "Duración", fmt="[h]:mm:ss", width=13, source="duration_s", conversion="Segundos ÷ 86.400", source_unit="s", shown_unit="duración Excel"),
        column("moving_time", "Tiempo en movimiento", fmt="[h]:mm:ss", width=19, source="moving_duration_s", conversion="Segundos ÷ 86.400", source_unit="s", shown_unit="duración Excel"),
        column("average_pace", "Ritmo medio (min/km)", fmt="[m]:ss", width=19, source="average_pace_s_per_km", conversion="Segundos ÷ 86.400", source_unit="s/km", shown_unit="min/km"),
        column("best_pace", "Mejor ritmo (min/km)", fmt="[m]:ss", width=19, source="best_pace_s_per_km", conversion="Segundos ÷ 86.400", source_unit="s/km", shown_unit="min/km"),
        column("average_hr", "FC media (lpm)", fmt="0", width=15, source="average_heart_rate_bpm", shown_unit="lpm"),
        column("maximum_hr", "FC máxima (lpm)", fmt="0", width=16, source="maximum_heart_rate_bpm", shown_unit="lpm"),
        column("average_power", "Potencia media (W)", fmt="0", width=18, source="average_power_w", shown_unit="W"),
        column("normalized_power", "Potencia normalizada (W)", fmt="0", width=23, source="normalized_power_w", shown_unit="W"),
        column("average_cadence", "Cadencia media", fmt="0", width=16, source="average_cadence_spm/average_cycling_cadence_rpm"),
        column("elevation_gain", "Desnivel positivo (m)", fmt="0", width=20, source="elevation_gain_m", shown_unit="m"),
        column("aerobic_effect", "Efecto aeróbico", fmt="0.0", width=16, source="aerobic_training_effect"),
        column("anaerobic_effect", "Efecto anaeróbico", fmt="0.0", width=18, source="anaerobic_training_effect"),
        column("benefit", "Beneficio principal", width=20, source="training_effect_label"),
        column("garmin_load", "Carga Garmin", fmt="0.0", width=14, source="training_load"),
        column("calories", "Calorías (kcal)", fmt="0", width=15, source="calories_kcal", shown_unit="kcal"),
        column("sweat", "Sudor estimado (ml)", fmt="0", width=19, source="estimated_sweat_loss_ml", shown_unit="ml"),
        column("rpe", "Esfuerzo percibido (1–10)", fmt="0.0", width=24, source="self_evaluation.perceived_exertion_1_10", shown_unit="1–10"),
        column("feeling", "Sensación", width=14, source="self_evaluation.feeling"),
        column("temperature", "Temperatura media (°C)", fmt="0.0", width=21, source="average_temperature_c", shown_unit="°C"),
        column("gear", "Equipamiento", width=34, source="gear_refs → gear", description="Nombre, marca y modelo asociados mediante la referencia privada de equipamiento."),
        column("activity_ref", "ID de actividad", width=18, hidden=True, source="activity_ref", description="Referencia privada estable para relacionar hojas."),
    ]

    interval_columns = [
        column("date", "Fecha", fmt="dd/mm/yyyy", width=12, source="activity.date"),
        column("activity", "Actividad", width=30, source="activity.name"),
        column("type", "Tipo de intervalo", width=20, source="interval_type"),
        column("level", "Nivel", width=18, source="xlsx.interval_level", description="Aclara si la fila resume toda la actividad, un bloque, un intervalo o una pausa."),
        column("count", "Número de intervalos", fmt="0", width=19, source="interval_count"),
        column("distance_km", "Distancia (km)", fmt="0.00", width=15, source="distance_m", conversion="Metros ÷ 1.000", source_unit="m", shown_unit="km"),
        column("duration", "Duración", fmt="[h]:mm:ss", width=13, source="duration_s", conversion="Segundos ÷ 86.400", source_unit="s", shown_unit="duración Excel"),
        column("moving_time", "Tiempo en movimiento", fmt="[h]:mm:ss", width=19, source="moving_duration_s", conversion="Segundos ÷ 86.400", source_unit="s", shown_unit="duración Excel"),
        column("average_pace", "Ritmo medio (min/km)", fmt="[m]:ss", width=20, source="average_pace_s_per_km", conversion="Segundos ÷ 86.400", source_unit="s/km", shown_unit="min/km"),
        column("moving_pace", "Ritmo en movimiento (min/km)", fmt="[m]:ss", width=25, source="moving_pace_s_per_km", conversion="Segundos ÷ 86.400", source_unit="s/km", shown_unit="min/km"),
        column("best_pace", "Mejor ritmo (min/km)", fmt="[m]:ss", width=20, source="best_pace_s_per_km", conversion="Segundos ÷ 86.400", source_unit="s/km", shown_unit="min/km"),
        column("average_hr", "FC media (lpm)", fmt="0", width=15, source="average_heart_rate_bpm", shown_unit="lpm"),
        column("maximum_hr", "FC máxima (lpm)", fmt="0", width=16, source="maximum_heart_rate_bpm", shown_unit="lpm"),
        column("average_power", "Potencia media (W)", fmt="0", width=18, source="average_power_w", shown_unit="W"),
        column("maximum_power", "Potencia máxima (W)", fmt="0", width=19, source="maximum_power_w", shown_unit="W"),
        column("cadence", "Cadencia media", fmt="0", width=16, source="average_cadence_spm"),
        column("calories", "Calorías (kcal)", fmt="0", width=15, source="calories_kcal", shown_unit="kcal"),
        column("elevation_gain", "Desnivel positivo (m)", fmt="0", width=20, source="elevation_gain_m", shown_unit="m"),
        column("elevation_loss", "Desnivel negativo (m)", fmt="0", width=20, source="elevation_loss_m", shown_unit="m"),
        column("gap", "Ritmo ajustado por pendiente (min/km)", fmt="[m]:ss", width=33, source="grade_adjusted_pace_s_per_km", conversion="Segundos ÷ 86.400", source_unit="s/km", shown_unit="min/km"),
        column("activity_ref", "ID de actividad", width=18, hidden=True, source="activity_ref"),
    ]

    lap_columns = [
        column("date", "Fecha", fmt="dd/mm/yyyy", width=12, source="activity.date"),
        column("activity", "Actividad", width=30, source="activity.name"),
        column("lap_index", "Número de vuelta", fmt="0", width=16, source="lap_index"),
        column("lap_type", "Tipo de vuelta", width=18, source="step_type/lap_type"),
        column("distance_m", "Distancia (m)", fmt="0", width=14, source="distance_m", shown_unit="m"),
        column("duration", "Duración", fmt="[h]:mm:ss", width=13, source="duration_s", conversion="Segundos ÷ 86.400", source_unit="s", shown_unit="duración Excel"),
        column("moving_time", "Tiempo en movimiento", fmt="[h]:mm:ss", width=19, source="moving_duration_s", conversion="Segundos ÷ 86.400", source_unit="s", shown_unit="duración Excel"),
        column("average_pace", "Ritmo medio (min/km)", fmt="[m]:ss", width=20, source="average_pace_s_per_km", conversion="Segundos ÷ 86.400", source_unit="s/km", shown_unit="min/km"),
        column("best_pace", "Mejor ritmo (min/km)", fmt="[m]:ss", width=20, source="best_pace_s_per_km", conversion="Segundos ÷ 86.400", source_unit="s/km", shown_unit="min/km"),
        column("average_hr", "FC media (lpm)", fmt="0", width=15, source="average_heart_rate_bpm", shown_unit="lpm"),
        column("maximum_hr", "FC máxima (lpm)", fmt="0", width=16, source="maximum_heart_rate_bpm", shown_unit="lpm"),
        column("average_power", "Potencia media (W)", fmt="0", width=18, source="average_power_w", shown_unit="W"),
        column("maximum_power", "Potencia máxima (W)", fmt="0", width=19, source="maximum_power_w", shown_unit="W"),
        column("normalized_power", "Potencia normalizada (W)", fmt="0", width=23, source="normalized_power_w", shown_unit="W"),
        column("cadence", "Cadencia", fmt="0", width=12, source="average_cadence_spm"),
        column("stride", "Longitud de zancada (cm)", fmt="0.0", width=23, source="average_stride_length_cm", shown_unit="cm"),
        column("ground_contact", "Contacto con el suelo (ms)", fmt="0", width=24, source="average_ground_contact_time_ms", shown_unit="ms"),
        column("vertical_oscillation", "Oscilación vertical (cm)", fmt="0.0", width=22, source="average_vertical_oscillation_cm", shown_unit="cm"),
        column("calories", "Calorías (kcal)", fmt="0", width=15, source="calories_kcal", shown_unit="kcal"),
        column("elevation_gain", "Desnivel positivo (m)", fmt="0", width=20, source="elevation_gain_m", shown_unit="m"),
        column("elevation_loss", "Desnivel negativo (m)", fmt="0", width=20, source="elevation_loss_m", shown_unit="m"),
        column("partial", "Vuelta parcial", width=14, source="partial_lap"),
        column("activity_ref", "ID de actividad", width=18, hidden=True, source="activity_ref"),
    ]

    zone_columns = [
        column("date", "Fecha", fmt="dd/mm/yyyy", width=12, source="activity.date"),
        column("activity", "Actividad", width=30, source="activity.name"),
        column("zone_type", "Tipo de zona", width=22, source="zone_type"),
        column("zone", "Zona", width=22, source="zone"),
        column("duration", "Duración", fmt="[h]:mm:ss", width=13, source="duration_s", conversion="Segundos ÷ 86.400", source_unit="s", shown_unit="duración Excel"),
        column("percentage", "Porcentaje", fmt="0.0%", width=14, source="duration_s/total_duration_s", shown_unit="%"),
        column("lower", "Límite inferior", fmt="0", width=15, source="low_boundary"),
        column("upper", "Límite superior", fmt="0", width=15, source="next_low_boundary"),
        column("unit", "Unidad", width=10, source="zone_type"),
        column("activity_ref", "ID de actividad", width=18, hidden=True, source="activity_ref"),
    ]

    daily_columns = [
        column("date", "Fecha", fmt="dd/mm/yyyy", width=12, source="date", group="Actividad diaria"),
        column("steps", "Pasos", fmt="0", width=11, source="steps", group="Actividad diaria"),
        column("distance_km", "Distancia diaria (km)", fmt="0.00", width=19, source="distance_m", conversion="Metros ÷ 1.000", source_unit="m", shown_unit="km", group="Actividad diaria"),
        column("active_calories", "Calorías activas", fmt="0", width=16, source="active_calories_kcal", shown_unit="kcal", group="Actividad diaria"),
        column("total_calories", "Calorías totales", fmt="0", width=16, source="total_calories_kcal", shown_unit="kcal", group="Actividad diaria"),
        column("moderate_minutes", "Intensidad moderada (min)", fmt="0", width=23, source="moderate_intensity_minutes", shown_unit="min", group="Actividad diaria"),
        column("vigorous_minutes", "Intensidad vigorosa (min)", fmt="0", width=23, source="vigorous_intensity_minutes", shown_unit="min", group="Actividad diaria"),
        column("resting_hr", "Pulso en reposo (lpm)", fmt="0", width=20, source="resting_heart_rate_bpm", shown_unit="lpm", group="Recuperación"),
        column("average_stress", "Estrés medio", fmt="0.0", width=14, source="average_stress", group="Recuperación"),
        column("maximum_stress", "Estrés máximo", fmt="0.0", width=15, source="maximum_stress", group="Recuperación"),
        column("body_battery_high", "Body Battery máximo", fmt="0", width=19, source="body_battery_high", group="Recuperación"),
        column("body_battery_low", "Body Battery mínimo", fmt="0", width=18, source="body_battery_low", group="Recuperación"),
        column("body_battery_charged", "Body Battery cargado", fmt="0", width=20, source="body_battery_charged", group="Recuperación"),
        column("body_battery_drained", "Body Battery consumido", fmt="0", width=21, source="body_battery_drained", group="Recuperación"),
        column("average_spo2", "SpO₂ media", fmt="0.0%", width=13, source="average_spo2_pct", conversion="Porcentaje ÷ 100", source_unit="0–100", shown_unit="%", group="Recuperación"),
        column("lowest_spo2", "SpO₂ mínima", fmt="0.0%", width=14, source="lowest_spo2_pct", conversion="Porcentaje ÷ 100", source_unit="0–100", shown_unit="%", group="Recuperación"),
        column("waking_respiration", "Respiración despierto", fmt="0.0", width=20, source="average_waking_respiration_brpm", shown_unit="resp/min", group="Recuperación"),
        column("sleep_available", "Sueño disponible", width=16, source="sleep.valid_sleep", group="Sueño"),
        column("sleep_start", "Inicio del sueño", fmt="dd/mm/yyyy hh:mm", width=19, source="sleep.sleep_start_local", group="Sueño"),
        column("sleep_end", "Final del sueño", fmt="dd/mm/yyyy hh:mm", width=19, source="sleep.sleep_end_local", group="Sueño"),
        column("total_sleep", "Sueño total", fmt="[h]:mm:ss", width=14, source="sleep.total_sleep_s", conversion="Segundos ÷ 86.400", source_unit="s", shown_unit="duración Excel", group="Sueño"),
        column("awake", "Tiempo despierto", fmt="[h]:mm:ss", width=16, source="sleep.awake_s", conversion="Segundos ÷ 86.400", source_unit="s", shown_unit="duración Excel", group="Sueño"),
        column("light_sleep", "Sueño ligero", fmt="[h]:mm:ss", width=14, source="sleep.light_sleep_s", conversion="Segundos ÷ 86.400", source_unit="s", shown_unit="duración Excel", group="Sueño"),
        column("deep_sleep", "Sueño profundo", fmt="[h]:mm:ss", width=15, source="sleep.deep_sleep_s", conversion="Segundos ÷ 86.400", source_unit="s", shown_unit="duración Excel", group="Sueño"),
        column("rem_sleep", "Sueño REM", fmt="[h]:mm:ss", width=13, source="sleep.rem_sleep_s", conversion="Segundos ÷ 86.400", source_unit="s", shown_unit="duración Excel", group="Sueño"),
        column("naps", "Siestas", fmt="[h]:mm:ss", width=12, source="sleep.nap_time_s", conversion="Segundos ÷ 86.400", source_unit="s", shown_unit="duración Excel", group="Sueño"),
        column("sleep_score", "Puntuación del sueño", fmt="0", width=19, source="sleep.sleep_score", group="Sueño"),
        column("sleep_rating", "Valoración del sueño", width=19, source="sleep.sleep_score_qualifier", group="Sueño"),
        column("sleep_hr", "Pulso durante el sueño", fmt="0", width=20, source="sleep.average_sleep_heart_rate_bpm", shown_unit="lpm", group="Sueño"),
        column("sleep_stress", "Estrés durante el sueño", fmt="0.0", width=21, source="sleep.average_sleep_stress", group="Sueño"),
        column("sleep_spo2", "SpO₂ durante el sueño", fmt="0.0%", width=20, source="sleep.average_sleep_spo2_pct", conversion="Porcentaje ÷ 100", source_unit="0–100", shown_unit="%", group="Sueño"),
        column("sleep_respiration", "Respiración durante el sueño", fmt="0.0", width=24, source="average_sleep_respiration_brpm", shown_unit="resp/min", group="Sueño"),
        column("hrv_average", "VFC nocturna media (ms)", fmt="0", width=22, source="hrv.overnight_average_ms", shown_unit="ms", group="VFC"),
        column("hrv_best", "Mejor media de 5 min (ms)", fmt="0", width=22, source="hrv.highest_five_min_average_ms", shown_unit="ms", group="VFC"),
        column("hrv_low", "Límite inferior VFC (ms)", fmt="0", width=21, source="hrv.baseline_balanced_low_ms", shown_unit="ms", group="VFC"),
        column("hrv_high", "Límite superior VFC (ms)", fmt="0", width=21, source="hrv.baseline_balanced_high_ms", shown_unit="ms", group="VFC"),
        column("hrv_status", "Estado de VFC", width=16, source="hrv.status", group="VFC"),
        column("hrv_week", "Media semanal VFC (ms)", fmt="0", width=20, source="hrv.weekly_average_ms", shown_unit="ms", group="VFC"),
    ]

    habit_columns = [
        column("date", "Fecha", fmt="dd/mm/yyyy", width=12, source="lifestyle_logs.date"),
        column("habit", "Hábito registrado", width=21, source="lifestyle_logs.behaviour"),
        column("status", "Estado", width=16, source="lifestyle_logs.status"),
        column("amount", "Cantidad", fmt="0.0", width=13, source="lifestyle_logs.amount"),
        column("note", "Nota", width=34, source="lifestyle_logs.note"),
    ]

    gear_columns = [
        column("name", "Nombre", width=28, source="gear_name", description="Nombre original del equipamiento."),
        column("model", "Marca y modelo", width=30, source="manufacturer/model", description="Fabricante y modelo confirmados por Garmin o por la persona."),
        column("type", "Tipo", width=15, source="type"),
        column("status", "Estado", width=13, source="status/retired"),
        column("total_km", "Distancia total Garmin (km)", fmt="0.00", width=25, source="total_distance_m", conversion="Metros ÷ 1.000", source_unit="m", shown_unit="km"),
        column("period_km", "Distancia en el periodo (km)", fmt="0.00", width=26, source="activities.distance_m", conversion="Suma por gear_ref y metros ÷ 1.000", source_unit="m", shown_unit="km"),
        column("activities", "Actividades en el periodo", fmt="0", width=23, source="activities.gear_refs"),
        column("last_date", "Última fecha de uso", fmt="dd/mm/yyyy", width=19, source="activities.date"),
        column("gear_ref", "ID interno", width=18, hidden=True, source="gear_ref"),
    ]

    journal_columns = [
        column("journal_date", "Fecha del diario", fmt="dd/mm/yyyy", width=17, source="date"),
        column("activity_date", "Fecha de la actividad", fmt="dd/mm/yyyy", width=20, source="activity.date"),
        column("activity", "Actividad", width=30, source="activity_ref → activity.name"),
        column("note", "Nota", width=45, source="note", description="Comentario que la persona decidió incluir."),
        column("planned_type", "Tipo de sesión previsto", width=22, source="intended_session_type"),
        column("warning", "Aviso", width=48, source="xlsx.date_mismatch", description="Advierte si las fechas del diario y de la actividad no coinciden."),
        column("activity_ref", "ID de actividad", width=18, hidden=True, source="activity_ref"),
    ]

    quality_columns = [
        column("metric", "Métrica", width=32, source="coverage.<metric>"),
        column("mean", "Media", fmt="0.0", width=12, source="mean"),
        column("median", "Mediana", fmt="0.0", width=12, source="median"),
        column("available", "Datos disponibles", fmt="0", width=18, source="available_days/available_activities"),
        column("expected", "Datos esperados", fmt="0", width=16, source="expected_days/expected_activities"),
        column("missing", "Datos ausentes", fmt="0", width=15, source="missing_days/missing_activities"),
        column("coverage", "Cobertura", fmt="0.0%", width=13, source="coverage_pct", conversion="Porcentaje ÷ 100", source_unit="0–100", shown_unit="%"),
        column("missing_periods", "Periodos sin datos", width=38, source="missing_date_ranges"),
        column("status", "Estado", width=12, source="coverage_pct"),
    ]

    # 1. INICIO
    start_sheet = workbook.create_sheet("INICIO")
    start_sheet.sheet_view.showGridLines = False
    start_sheet.column_dimensions["A"].width = 34
    start_sheet.column_dimensions["B"].width = 58
    title = WriteOnlyCell(start_sheet, value="Exportación de entrenamiento Garmin")
    title.font = title_font
    start_sheet.append([title])
    subtitle = WriteOnlyCell(start_sheet, value="Informe deportivo preparado para corredores y entrenadores")
    subtitle.font = secondary_font
    start_sheet.append([subtitle])
    start_sheet.append([])

    event = race_context.get("event") or {}
    goal = race_context.get("goal") or {}
    availability = race_context.get("availability") or {}
    period = summary.get("period") or {}
    export_status = metadata.get("export_status", "completed")
    start_fields = [
        ("Fecha y hora de exportación", excel_datetime(metadata.get("exported_at")), "dd/mm/yyyy hh:mm", "export_metadata.exported_at", "Fecha y hora reales de Excel", "fecha y hora"),
        ("Inicio del periodo analizado", excel_date(period.get("start_date") or metadata.get("start_date")), "dd/mm/yyyy", "period_summary.period.start_date", "Fecha real de Excel", "fecha"),
        ("Fin del periodo analizado", excel_date(period.get("end_date") or metadata.get("end_date")), "dd/mm/yyyy", "period_summary.period.end_date", "Fecha real de Excel", "fecha"),
        ("Nombre de la carrera", user_text(event.get("label")), "General", "race_context.event.label", "Sin conversión", "texto"),
        ("Fecha de la carrera", excel_date(event.get("race_date")), "dd/mm/yyyy", "race_context.event.race_date", "Fecha real de Excel", "fecha"),
        ("Distancia de la carrera (km)", excel_km(event.get("distance_m")), "0.00", "race_context.event.distance_m", "Metros ÷ 1.000", "km"),
        ("Objetivo de tiempo", excel_time(goal.get("target_time_s")), "[h]:mm:ss", "race_context.goal.target_time_s", "Segundos ÷ 86.400", "duración"),
        ("Ritmo objetivo", excel_time(goal.get("target_pace_s_per_km")), '[m]:ss "min/km"', "race_context.goal.target_pace_s_per_km", "Segundos ÷ 86.400", "min/km"),
        ("Días restantes", event.get("days_remaining"), "0", "race_context.event.days_remaining", "Sin conversión", "días"),
        ("Semanas restantes", event.get("weeks_remaining"), "0.0", "race_context.event.weeks_remaining", "Sin conversión", "semanas"),
        ("Días disponibles para entrenar", availability.get("available_days_per_week"), "0", "race_context.availability.available_days_per_week", "Sin conversión", "días por semana"),
        ("Día previsto para la tirada larga", user_text(availability.get("long_run_day")), "General", "race_context.availability.long_run_day", "Sin conversión", "día de la semana"),
        ("Restricciones indicadas", user_text(availability.get("restrictions")), "General", "race_context.availability.restrictions", "Sin conversión", "texto"),
        ("Reloj principal", user_text(profile.get("primary_watch")), "General", "profile.primary_watch", "Sin conversión", "texto"),
        ("Estado de la exportación", "Completa" if export_status == "completed" else "Parcial", "General", "export_metadata.export_status", "Traducción del estado", "estado"),
    ]
    for label, value, fmt, source, conversion, unit in start_fields:
        label_cell = WriteOnlyCell(start_sheet, value=label)
        label_cell.font = Font(name="Aptos", size=10, bold=True, color=colors["navy"])
        label_cell.fill = light_fill
        label_cell.border = border
        value_cell = WriteOnlyCell(start_sheet, value=_excel_safe(value))
        value_cell.font = normal_font
        value_cell.number_format = fmt
        value_cell.alignment = Alignment(wrap_text=True, vertical="top")
        value_cell.border = border
        start_sheet.append([label_cell, value_cell])
        add_mapping_entry("INICIO", label, source, conversion, unit, f"Dato de portada: {label.lower()}.")
    start_sheet.append([])
    notice_title = WriteOnlyCell(start_sheet, value="Antes de compartir este archivo")
    notice_title.font = section_font
    start_sheet.append([notice_title])
    notices = [
        "Este archivo contiene datos privados de salud, ubicación y entrenamiento.",
        "Las celdas vacías significan dato ausente; nunca se interpretan como cero.",
        "ExportaGarmin no envía este archivo automáticamente a ninguna inteligencia artificial.",
        f"Modelo de datos: {metadata.get('schema_version', SCHEMA_VERSION)} · Presentación Excel: {XLSX_PRESENTATION_VERSION}.",
    ]
    for notice in notices:
        cell = WriteOnlyCell(start_sheet, value=notice)
        cell.font = normal_font
        cell.fill = warning_fill if "privados" in notice else light_fill
        cell.alignment = Alignment(wrap_text=True)
        start_sheet.append([cell])

    # 2. RESUMEN: panel sin fórmulas; los cálculos proceden del modelo aprobado.
    summary_sheet = workbook.create_sheet("RESUMEN")
    summary_sheet.sheet_view.showGridLines = False
    summary_sheet.freeze_panes = "A3"
    for letter, width in {"A": 36, "B": 15, "C": 3, "D": 31, "E": 15, "F": 3, "G": 3}.items():
        summary_sheet.column_dimensions[letter].width = width
    for letter in ("H", "I", "J", "K", "L", "M", "N"):
        summary_sheet.column_dimensions[letter].hidden = True

    kpis = [
        ("Kilómetros de carrera", excel_km(nested(coach_summary, "running.distance_m")), "0.00", "km"),
        ("Sesiones de carrera", nested(coach_summary, "running.sessions"), "0", "sesiones"),
        ("Tiempo total corriendo", excel_time(nested(coach_summary, "running.duration_s")), "[h]:mm:ss", "duración"),
        ("Tirada más larga", excel_km(nested(coach_summary, "running.longest_run_distance_m")), "0.00", "km"),
        ("Días entrenados", nested(coach_summary, "training.days_with_any_training"), "0", "días"),
        ("Sesiones de ciclismo", nested(coach_summary, "cycling.sessions"), "0", "sesiones"),
        ("Sesiones de fuerza", nested(coach_summary, "strength.sessions"), "0", "sesiones"),
        ("Carga Garmin total", nested(coach_summary, "load.garmin_training_load_total"), "0.0", "puntos Garmin"),
        ("Esfuerzo percibido medio", nested(coach_summary, "load.average_rpe_1_10"), "0.0", "1–10"),
        ("Cobertura de autoevaluación", excel_pct(nested(coach_summary, "load.self_evaluation_coverage_pct")), "0.0%", "%"),
        ("Cobertura de sueño", excel_pct(nested(visible_quality, "coverage.sleep_duration_s.coverage_pct")), "0.0%", "%"),
        ("Semanas restantes para la carrera", event.get("weeks_remaining"), "0.0", "semanas"),
    ]
    comparison = compare_complete_week_blocks(coach_weeks)
    comparison_specs = [
        ("Kilómetros corriendo", "running_distance_m", excel_km, "0.00"),
        ("Tiempo corriendo", "running_duration_s", excel_time, "[h]:mm:ss"),
        ("Sesiones", "running_sessions", lambda value: value, "0"),
        ("Tirada larga", "longest_run_distance_m", excel_km, "0.00"),
        ("Carga Garmin", "garmin_training_load_total", lambda value: value, "0.0"),
        ("Carga por percepción de esfuerzo", "session_rpe_load_total", lambda value: value, "0.0"),
    ]
    grid_rows = max(42, len(week_rows) + 3)
    grid = [[None] * 14 for _ in range(grid_rows)]
    formats = {}
    grid[0][0] = "Resumen de la preparación"
    for index in range(6):
        left = kpis[index * 2]
        right = kpis[index * 2 + 1]
        row_index = index + 2
        grid[row_index][0], grid[row_index][1] = left[0], left[1]
        grid[row_index][3], grid[row_index][4] = right[0], right[1]
        formats[(row_index, 1)] = left[2]
        formats[(row_index, 4)] = right[2]
    if comparison.get("status") == "available":
        block_size = comparison["block_size"]
        grid[10][0] = (
            f"Comparación de {block_size} semanas completas con las "
            f"{block_size} anteriores"
        )
    else:
        grid[10][0] = "Comparación de semanas completas"
    comparison_columns = (0, 1, 3, 4)
    for column_index, header in zip(
        comparison_columns,
        (
            "Métrica",
            f"{comparison.get('block_size', 0)} semanas anteriores",
            f"{comparison.get('block_size', 0)} semanas recientes",
            "Cambio",
        ),
    ):
        grid[11][column_index] = header
    comparison_metrics = comparison.get("metrics", {}) if comparison.get("status") == "available" else {}
    for row_offset, (label, metric, converter, fmt) in enumerate(comparison_specs, 12):
        values = comparison_metrics.get(metric, {})
        grid[row_offset][0] = label
        grid[row_offset][1] = converter(values.get("previous_block_total"))
        grid[row_offset][3] = converter(values.get("recent_block_total"))
        grid[row_offset][4] = excel_pct(values.get("percentage_change"))
        formats[(row_offset, 1)] = fmt
        formats[(row_offset, 3)] = fmt
        formats[(row_offset, 4)] = "0.0%"
    if not comparison_metrics:
        grid[12][0] = (
            "Hay "
            f"{comparison.get('complete_weeks_available', 0)} semanas completas; "
            f"se necesitan al menos {comparison.get('minimum_required', 2)}."
        )

    grid[0][7], grid[0][8], grid[0][9] = "Semana", "Carrera (km)", "Tirada larga (km)"
    grid[0][13] = "Inicio de semana"
    for index, week in enumerate(week_rows, 1):
        grid[index][7] = week.get("week")
        grid[index][8] = week.get("run_km")
        grid[index][9] = week.get("longest_km")
        grid[index][13] = week.get("from")
        formats[(index, 13)] = "dd/mm/yyyy"
    grid[0][11], grid[0][12] = "Deporte", "Tiempo"
    sport_distribution = [
        ("Carrera", excel_time(nested(coach_summary, "running.duration_s"))),
        ("Ciclismo", excel_time(nested(coach_summary, "cycling.duration_s"))),
        ("Fuerza", excel_time(nested(coach_summary, "strength.duration_s"))),
        ("Otros", excel_time(nested(coach_summary, "other.duration_s"))),
    ]
    for index, (sport_name, duration) in enumerate(sport_distribution, 1):
        grid[index][11], grid[index][12] = sport_name, duration

    for row_index, values in enumerate(grid):
        cells = []
        for column_index, value in enumerate(values):
            cell = WriteOnlyCell(summary_sheet, value=_excel_safe(value))
            cell.font = normal_font
            cell.number_format = formats.get((row_index, column_index), "General")
            if row_index == 0 and column_index == 0:
                cell.font = title_font
            elif row_index in (10,) and column_index == 0:
                cell.font = section_font
            elif row_index in (2, 3, 4, 5, 6, 7) and column_index in (0, 3):
                cell.font = Font(name="Aptos", size=10, bold=True, color=colors["navy"])
                cell.fill = light_fill
            elif row_index == 11 and column_index in comparison_columns:
                cell.font = header_font
                cell.fill = header_fill
            cell.alignment = Alignment(
                wrap_text=not (row_index == 0 and column_index == 0),
                vertical="center",
            )
            cells.append(cell)
        summary_sheet.append(cells)
    if comparison_metrics:
        summary_sheet.conditional_formatting.add(
            "E13:E18",
            ColorScaleRule(
                start_type="min", start_color="FCE5CD",
                mid_type="percentile", mid_value=50, mid_color="FFFFFF",
                end_type="max", end_color="D9EAF7",
            ),
        )
    if week_rows:
        categories = Reference(summary_sheet, min_col=14, min_row=2, max_row=len(week_rows) + 1)
        km_data = Reference(summary_sheet, min_col=9, min_row=1, max_row=len(week_rows) + 1)
        long_data = Reference(summary_sheet, min_col=10, min_row=1, max_row=len(week_rows) + 1)
        km_chart = BarChart()
        km_chart.x_axis = DateAxis(axId=500, crossAx=100)
        km_chart.y_axis.crossAx = 500
        km_chart.title = "Kilómetros de carrera por semana"
        km_chart.y_axis.numFmt = "0.0"
        km_chart.y_axis.tickLblPos = "nextTo"
        km_chart.y_axis.delete = False
        km_chart.x_axis.axPos = "b"
        km_chart.x_axis.tickLblPos = "nextTo"
        km_chart.x_axis.delete = False
        km_chart.x_axis.number_format = "dd/mm"
        km_chart.x_axis.majorTimeUnit = "days"
        km_chart.x_axis.majorUnit = 14
        km_chart.add_data(km_data, titles_from_data=True)
        km_chart.set_categories(categories)
        km_chart.height, km_chart.width = 7, 12
        km_chart.visible_cells_only = False
        km_chart.varyColors = False
        km_chart.legend = None
        km_chart.series[0].graphicalProperties.solidFill = colors["blue"]
        summary_sheet.add_chart(km_chart, "F2")
        long_chart = LineChart()
        long_chart.x_axis = DateAxis(axId=500, crossAx=100)
        long_chart.y_axis.crossAx = 500
        long_chart.title = "Evolución de la tirada más larga (km)"
        long_chart.y_axis.numFmt = "0.0"
        long_chart.y_axis.tickLblPos = "nextTo"
        long_chart.y_axis.delete = False
        long_chart.x_axis.axPos = "b"
        long_chart.x_axis.tickLblPos = "nextTo"
        long_chart.x_axis.delete = False
        long_chart.x_axis.number_format = "dd/mm"
        long_chart.x_axis.majorTimeUnit = "days"
        long_chart.x_axis.majorUnit = 14
        long_chart.add_data(long_data, titles_from_data=True)
        long_chart.set_categories(categories)
        long_chart.height, long_chart.width = 7, 12
        long_chart.visible_cells_only = False
        long_chart.varyColors = False
        long_chart.legend = None
        long_chart.series[0].graphicalProperties.line.solidFill = colors["blue"]
        summary_sheet.add_chart(long_chart, "F17")
    if any(value for _, value in sport_distribution):
        distribution_data = Reference(summary_sheet, min_col=13, min_row=1, max_row=5)
        distribution_labels = Reference(summary_sheet, min_col=12, min_row=2, max_row=5)
        sport_chart = PieChart()
        sport_chart.title = "Distribución del tiempo por deporte"
        sport_chart.add_data(distribution_data, titles_from_data=True)
        sport_chart.set_categories(distribution_labels)
        sport_chart.height, sport_chart.width = 7, 10
        sport_chart.visible_cells_only = False
        sport_chart.varyColors = True
        summary_sheet.add_chart(sport_chart, "F32")

    for label, _, _, unit in kpis:
        add_mapping_entry("RESUMEN", label, "period_summary/race_context/data_quality", "Cálculo semántico aprobado", unit, f"Indicador destacado: {label.lower()}.")

    # Hojas tabulares visibles en el orden solicitado.
    add_table_sheet(
        "SEMANAS",
        week_rows,
        week_columns,
        freeze="D2",
        primary_keys={
            "week", "status", "sessions", "training_days", "run_km",
            "run_time", "longest_km", "garmin_load", "average_rpe",
            "sleep", "hrv", "resting_hr",
        },
    )
    add_table_sheet(
        "ACTIVIDADES",
        activity_rows,
        activity_columns,
        freeze="C2",
        primary_keys={
            "date", "name", "sport", "session_type", "record_status",
            "distance_km", "duration", "average_pace", "average_hr",
            "garmin_load", "rpe", "feeling", "gear",
        },
    )
    add_table_sheet("INTERVALOS", interval_rows, interval_columns, freeze="C2")
    add_table_sheet("VUELTAS", lap_rows, lap_columns, freeze="C2")
    zone_sheet = add_table_sheet("ZONAS", zone_rows, zone_columns, freeze="C2")
    if zone_sheet and zone_rows:
        zone_sheet.conditional_formatting.add(
            f"F2:F{len(zone_rows) + 1}",
            ColorScaleRule(
                start_type="min", start_color="FFFFFF",
                end_type="max", end_color="D9EAF7",
            ),
        )
    add_table_sheet(
        "SALUD DIARIA",
        daily_rows,
        daily_columns,
        freeze="B2",
        primary_keys={
            "date", "steps", "resting_hr", "average_stress",
            "body_battery_high", "total_sleep", "sleep_score", "hrv_average",
        },
    )
    add_table_sheet("HÁBITOS", habit_rows, habit_columns, freeze="B2")

    # MEDIDAS contiene dos tablas independientes y un único gráfico de peso.
    if composition_rows or pressure_rows:
        measures_sheet = workbook.create_sheet("MEDIDAS")
        measures_sheet.sheet_view.showGridLines = False
        comp_columns = [
            column("date", "Fecha", fmt="dd/mm/yyyy", width=19, source="body_composition.date"),
            column("weight", "Peso (kg)", fmt="0.00", width=23, source="weight_kg", shown_unit="kg"),
            column("bmi", "IMC", fmt="0.0", width=24, source="bmi"),
            column("fat", "Grasa corporal", fmt="0.0%", width=16, source="body_fat_pct", conversion="Porcentaje ÷ 100", source_unit="0–100", shown_unit="%"),
            column("water", "Agua corporal", fmt="0.0%", width=16, source="body_water_pct", conversion="Porcentaje ÷ 100", source_unit="0–100", shown_unit="%"),
            column("muscle", "Masa muscular (kg)", fmt="0.00", width=19, source="muscle_mass_kg", shown_unit="kg"),
            column("bone", "Masa ósea (kg)", fmt="0.00", width=16, source="bone_mass_kg", shown_unit="kg"),
        ]
        pressure_columns = [
            column("datetime", "Fecha y hora", fmt="dd/mm/yyyy hh:mm", width=19, source="timestamp"),
            column("systolic", "Presión sistólica (mmHg)", fmt="0", width=23, source="systolic_mmhg", shown_unit="mmHg"),
            column("diastolic", "Presión diastólica (mmHg)", fmt="0", width=24, source="diastolic_mmhg", shown_unit="mmHg"),
            column("pulse", "Pulso (lpm)", fmt="0", width=14, source="pulse_bpm", shown_unit="lpm"),
        ]
        current_row = 1
        if composition_rows:
            current_row = append_table(measures_sheet, composition_rows, comp_columns, "Composicion", current_row)
            register_mapping("MEDIDAS", comp_columns)
            weight_chart = LineChart()
            weight_chart.x_axis = DateAxis(axId=500, crossAx=100)
            weight_chart.y_axis.crossAx = 500
            weight_chart.title = "Evolución del peso (kg)"
            weight_chart.y_axis.numFmt = "0.00"
            weight_chart.y_axis.tickLblPos = "nextTo"
            weight_chart.y_axis.delete = False
            weight_chart.x_axis.axPos = "b"
            weight_chart.x_axis.tickLblPos = "nextTo"
            weight_chart.x_axis.delete = False
            weight_chart.x_axis.number_format = "dd/mm"
            weight_chart.x_axis.majorTimeUnit = "days"
            weight_chart.x_axis.majorUnit = 18
            weight_chart.add_data(
                Reference(measures_sheet, min_col=2, min_row=1, max_row=len(composition_rows) + 1),
                titles_from_data=True,
            )
            weight_chart.set_categories(
                Reference(measures_sheet, min_col=1, min_row=2, max_row=len(composition_rows) + 1)
            )
            weight_chart.height, weight_chart.width = 7, 11
            weight_chart.varyColors = False
            weight_chart.legend = None
            weight_chart.series[0].graphicalProperties.line.solidFill = colors["blue"]
            measures_sheet.add_chart(weight_chart, "I2")
        if pressure_rows:
            for _ in range(3):
                measures_sheet.append([])
            current_row += 4
            append_table(measures_sheet, pressure_rows, pressure_columns, "PresionArterial", current_row)
            register_mapping("MEDIDAS", pressure_columns)
        measures_sheet.freeze_panes = "A2"

    add_table_sheet("EQUIPAMIENTO", gear_rows, gear_columns, freeze="B2")
    add_table_sheet("DIARIO", journal_rows, journal_columns, freeze="B2")

    # CALIDAD DATOS: tabla principal y secciones narrativas sin rutas internas.
    quality_sheet = workbook.create_sheet("CALIDAD DATOS")
    quality_sheet.freeze_panes = "A2"
    quality_sheet.column_dimensions["A"].width = 34
    quality_sheet.column_dimensions["B"].width = 14
    quality_sheet.column_dimensions["C"].width = 14
    quality_sheet.column_dimensions["D"].width = 18
    quality_sheet.column_dimensions["E"].width = 17
    quality_sheet.column_dimensions["F"].width = 16
    quality_sheet.column_dimensions["G"].width = 14
    quality_sheet.column_dimensions["H"].width = 42
    quality_sheet.column_dimensions["I"].width = 13
    current_quality_row = append_table(quality_sheet, quality_rows or [{
        "metric": "No hay métricas de cobertura disponibles",
    }], quality_columns, "CalidadDatos", 1)
    register_mapping("CALIDAD DATOS", quality_columns)
    if quality_rows:
        quality_range = f"G2:G{len(quality_rows) + 1}"
        quality_sheet.conditional_formatting.add(
            quality_range,
            CellIsRule(operator="greaterThanOrEqual", formula=["0.8"], fill=good_fill),
        )
        quality_sheet.conditional_formatting.add(
            quality_range,
            CellIsRule(operator="between", formula=["0.5", "0.799999"], fill=warning_fill),
        )
        quality_sheet.conditional_formatting.add(
            quality_range,
            CellIsRule(
                operator="lessThan",
                formula=["0.5"],
                fill=PatternFill("solid", fgColor=colors["red"]),
            ),
        )

    def append_quality_section(title_text, values):
        nonlocal current_quality_row
        quality_sheet.append([])
        current_quality_row += 1
        heading = WriteOnlyCell(quality_sheet, value=title_text)
        heading.font = section_font
        heading.fill = light_fill
        quality_sheet.append([heading])
        current_quality_row += 1
        if not values:
            values = ["No se registraron elementos en esta sección."]
        for value in values:
            text_value = narrative(value)
            cell = WriteOnlyCell(quality_sheet, value=_excel_safe(text_value))
            cell.font = normal_font
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            quality_sheet.append([cell])
            current_quality_row += 1

    append_quality_section("Avisos", [
        *(visible_quality.get("warnings") or []),
        *(visible_quality.get("issues") or []),
        *(
            [
                f"Se conservaron {len(microactivities)} registros muy breves en "
                "ACTIVIDADES, pero no cuentan en sesiones, días entrenados, "
                "carga, comparaciones ni gráficos. Fechas: "
                + ", ".join(
                    sorted({
                        item.get("date")
                        for item in microactivities
                        if item.get("date")
                    })
                )
                + "."
            ]
            if microactivities else []
        ),
    ])
    append_quality_section("Limitaciones", visible_quality.get("limitations") or [])
    append_quality_section("Transformaciones aplicadas", visible_quality.get("transformations") or [])
    privacy = visible_quality.get("privacy") or {}
    privacy_lines = [
        "La identidad y los identificadores personales se ocultan automáticamente.",
        "Las coordenadas y los datos deportivos se conservan para el análisis.",
        "No se incluyen credenciales, tokens ni cachés privadas.",
        f"Política aplicada: {translate('privacy', privacy.get('mode')) or 'Automática'}.",
    ]
    append_quality_section("Privacidad", privacy_lines)
    append_quality_section("Deduplicaciones", visible_quality.get("deduplication") or [])
    append_quality_section("Series temporales en Excel", [
        (
            f"Se incluyeron {available_series_samples} muestras en la hoja "
            "oculta DATOS POR SEGUNDO."
            if available_series_samples and not omit_series
            else (
                "Se omitieron del XLSX "
                f"{available_series_samples:,} muestras ".replace(",", ".")
                + "porque superan el límite de "
                + f"{XLSX_MAX_ACTIVITY_SERIES_SAMPLES:,}. ".replace(",", ".")
                +
                "El TXT conserva las series completas."
                if omit_series
                else "No había series temporales disponibles."
            )
        )
    ])

    # AYUDA se construye desde el mapeo real de las hojas visibles.
    help_rows = []
    seen_help = set()
    for item in mapping_rows:
        key = (item["sheet"], item["visible_name"])
        if key in seen_help:
            continue
        seen_help.add(key)
        help_rows.append({
            "sheet": item["sheet"],
            "field": item["visible_name"],
            "meaning": item["description"],
            "unit": item["shown_unit"],
            "source": (
                "Contexto indicado por la persona"
                if str(item["internal_field"]).startswith("race_context")
                else (
                    "Anotación indicada por la persona"
                    if str(item["internal_field"]).startswith("journal")
                    else (
                        "Cálculo de ExportaGarmin"
                        if any(
                            source_name in str(item["internal_field"])
                            for source_name in (
                                "period_summary", "race_analysis", "data_quality"
                            )
                        )
                        else "Garmin Connect"
                    )
                )
            ),
            "notes": item["conversion"],
        })
    help_columns = [
        column("sheet", "Hoja", width=20, source="technical_mapping.sheet"),
        column("field", "Campo visible", width=32, source="technical_mapping.visible_name"),
        column("meaning", "Qué significa", width=55, source="technical_mapping.description"),
        column("unit", "Unidad", width=18, source="technical_mapping.shown_unit"),
        column("source", "Procedencia", width=28, source="technical_mapping.source"),
        column("notes", "Observaciones", width=38, source="technical_mapping.conversion"),
    ]
    add_table_sheet("AYUDA", help_rows, help_columns, freeze="B2")

    # Hojas técnicas ocultas. Conservan valores originales y precisión sin
    # contaminar las vistas de usuario con rutas internas o JSON.
    technical_activity_rows = []

    def scalar_rows(value, path_prefix=""):
        rows = []
        if isinstance(value, dict):
            if not value:
                rows.append((path_prefix, None, True))
            for key, nested_value in value.items():
                if key == "activity_series":
                    continue
                path_name = f"{path_prefix}.{key}" if path_prefix else str(key)
                rows.extend(scalar_rows(nested_value, path_name))
        elif isinstance(value, list):
            if not value:
                rows.append((path_prefix, None, True))
            for index, nested_value in enumerate(value):
                rows.extend(scalar_rows(nested_value, f"{path_prefix}[{index}]"))
        else:
            rows.append((path_prefix, value, value is None))
        return rows

    for activity in activities:
        reference = activity.get("activity_ref")
        for path_name, original_value, is_null in scalar_rows(activity):
            technical_activity_rows.append({
                "activity_ref": reference,
                "field": path_name,
                "value": original_value,
                "is_null": is_null,
            })
    technical_activity_columns = [
        column("activity_ref", "Referencia privada", width=21),
        column("field", "Campo interno", width=48),
        column("value", "Valor original", width=38),
        column("is_null", "Era nulo", width=12),
    ]
    add_table_sheet(
        "TÉCNICO - ACTIVIDADES",
        technical_activity_rows,
        technical_activity_columns,
        hidden=True,
    )

    technical_model_rows = []
    for section_name, section_value in model.items():
        if section_name == "activities":
            continue
        for path_name, original_value, is_null in scalar_rows(section_value, section_name):
            technical_model_rows.append({
                "section": section_name,
                "field": path_name,
                "value": original_value,
                "is_null": is_null,
            })
    technical_model_columns = [
        column("section", "Sección interna", width=25),
        column("field", "Campo interno", width=55),
        column("value", "Valor original", width=42),
        column("is_null", "Era nulo", width=12),
    ]
    add_table_sheet(
        "TÉCNICO - MODELO",
        technical_model_rows,
        technical_model_columns,
        hidden=True,
    )

    technical_mapping_columns = [
        column("sheet", "Hoja visible", width=22),
        column("visible_name", "Nombre visible en español", width=36),
        column("internal_field", "Campo interno original", width=52),
        column("conversion", "Conversión aplicada", width=34),
        column("source_unit", "Unidad de origen", width=18),
        column("shown_unit", "Unidad mostrada", width=19),
    ]
    add_table_sheet(
        "TÉCNICO - MAPEO",
        mapping_rows,
        technical_mapping_columns,
        hidden=True,
    )

    if descriptor_rows and not omit_series:
        descriptor_columns = [
            column("activity_ref", "Referencia privada", width=21),
            column("descriptor_index", "Índice", fmt="0", width=10),
            column("column_name", "Columna técnica", width=30),
            column("field", "Campo normalizado", width=32),
            column("source_field", "Campo Garmin", width=32),
            column("source_unit", "Unidad Garmin", width=18),
            column("source_factor", "Factor", fmt="0.########", width=12),
        ]
        add_table_sheet(
            "TÉCNICO - SERIES",
            descriptor_rows,
            descriptor_columns,
            hidden=True,
        )
        sample_rows = []
        global_positions = {name: index for index, name in enumerate(series_columns)}
        for prepared in prepared_series:
            for sample_index, sample in enumerate(prepared.get("samples", [])):
                if not isinstance(sample, list):
                    continue
                row = {
                    "activity_ref": prepared.get("activity_ref"),
                    "sample_index": sample_index,
                }
                for local_index, column_name in enumerate(prepared.get("columns", [])):
                    if local_index < len(sample) and column_name in global_positions:
                        row[column_name] = sample[local_index]
                sample_rows.append(row)
        sample_columns = [
            column("activity_ref", "Referencia privada", width=21),
            column("sample_index", "Índice de muestra", fmt="0", width=18),
            *[
                column(name, humanize(name), fmt="0.########", width=18)
                for name in series_columns
            ],
        ]
        add_table_sheet(
            "DATOS POR SEGUNDO",
            sample_rows,
            sample_columns,
            hidden=True,
        )

    temporary = path.with_name(f".{path.name}.{secrets.token_hex(4)}.tmp")
    try:
        workbook.save(temporary)
        os.replace(temporary, path)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise
